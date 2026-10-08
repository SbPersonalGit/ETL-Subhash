"""
Unit tests for object hierarchy workbook generation.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path
from types import SimpleNamespace

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook
import requests

from tests._runner import ensure_project_root_on_path
ensure_project_root_on_path()

from core.object_hierarchy import (
    MappingRow,
    build_cycle_groups,
    build_hierarchy_analytics,
    build_hierarchy_groups,
    build_parent_id_mappings,
    generate_object_hierarchy_workbook,
)


class _FakeConnector:
    def __init__(self, describes: dict[str, dict]):
        self._describes = describes

    def describe_object(self, object_name: str) -> dict:
        return self._describes[object_name]


class _ActivityConnector(_FakeConnector):
    def describe_object(self, object_name: str) -> dict:
        if object_name == "Activity":
            raise requests.HTTPError(
                "not found", response=SimpleNamespace(status_code=404)
            )
        return super().describe_object(object_name)


def test_build_parent_id_mappings_filters_to_in_scope_objects():
    connector = _FakeConnector(
        {
            "Booking__c": {
                "fields": [
                    {"name": "Account__c", "type": "reference", "referenceTo": ["Account"]},
                    {"name": "OwnerId", "type": "reference", "referenceTo": ["User"]},
                    {"name": "Name", "type": "string"},
                ]
            },
            "Account": {"fields": []},
        }
    )

    rows = build_parent_id_mappings(connector, ["Booking__c", "Account"])

    assert rows == [MappingRow("Booking__c", "Account", "Account__c")]


def test_build_parent_id_mappings_excludes_user_lookups():
    connector = _FakeConnector(
        {
            "Passport__c": {
                "fields": [
                    {"name": "CreatedById", "type": "reference", "referenceTo": ["User"]},
                    {"name": "Account__c", "type": "reference", "referenceTo": ["Account"]},
                ]
            },
            "Account": {"fields": []},
            "User": {
                "fields": [
                    {"name": "AccountId", "type": "reference", "referenceTo": ["Account"]}
                ]
            },
        }
    )

    rows = build_parent_id_mappings(connector, ["Passport__c", "Account", "User"])

    assert rows == [
        MappingRow("Passport__c", "Account", "Account__c"),
        MappingRow("User", "Account", "AccountId"),
    ]


def test_build_parent_id_mappings_excludes_non_rest_parent_objects():
    connector = _ActivityConnector(
        {
            "Account": {
                "fields": [
                    {"name": "ActivityId", "type": "reference", "referenceTo": ["Activity"]}
                ]
            }
        }
    )
    skipped: list[dict[str, str]] = []

    rows = build_parent_id_mappings(
        connector, ["Account", "Activity"], skipped_objects=skipped
    )

    assert rows == []
    assert skipped == [
        {
            "object": "Activity",
            "reason": "Object is listed globally but has no REST describe resource.",
        }
    ]


def test_build_parent_id_mappings_excludes_noncreateable_audit_owner_and_polymorphic_fields():
    connector = _FakeConnector(
        {
            "Work__c": {
                "fields": [
                    {
                        "name": "Account__c",
                        "type": "reference",
                        "createable": True,
                        "referenceTo": ["Account"],
                    },
                    {
                        "name": "ReadonlyAccount__c",
                        "type": "reference",
                        "createable": False,
                        "referenceTo": ["Account"],
                    },
                    {
                        "name": "CreatedById",
                        "type": "reference",
                        "createable": True,
                        "referenceTo": ["User"],
                    },
                    {
                        "name": "OwnerId",
                        "type": "reference",
                        "createable": True,
                        "referenceTo": ["User", "Group"],
                    },
                    {
                        "name": "RelatedToId",
                        "type": "reference",
                        "createable": True,
                        "polymorphicForeignKey": True,
                        "referenceTo": ["Account"],
                    },
                    {
                        "name": "WhatId",
                        "type": "reference",
                        "createable": True,
                        "referenceTo": ["Account", "Opportunity"],
                    },
                ]
            },
            "Account": {"fields": []},
            "Opportunity": {"fields": []},
            "User": {"fields": []},
            "Group": {"fields": []},
        }
    )

    lookup_details: dict[str, dict[str, list[str]]] = {}
    rows = build_parent_id_mappings(
        connector,
        ["Work__c", "Account", "Opportunity", "User", "Group"],
        lookup_details=lookup_details,
    )

    assert rows == [MappingRow("Work__c", "Account", "Account__c")]
    assert lookup_details["Work__c"]["record_lookups"] == [
        "Account__c → Account"
    ]
    assert lookup_details["Work__c"]["other_lookups"] == [
        "ReadonlyAccount__c → Account [non-createable]",
        "CreatedById → User [audit/ownership]",
        "OwnerId → User, Group [audit/ownership]",
        "RelatedToId → Account [polymorphic]",
        "WhatId → Account, Opportunity [polymorphic]",
    ]


def test_build_cycle_groups_detects_closed_loop():
    groups = build_cycle_groups(
        [
            MappingRow("A__c", "B__c", "B__c"),
            MappingRow("B__c", "A__c", "A__c"),
            MappingRow("C__c", "B__c", "B__c"),
        ]
    )

    assert groups == [["A__c", "B__c"]]


def test_build_hierarchy_groups_orders_parent_groups_before_children():
    groups = build_hierarchy_groups(
        [
            MappingRow("Brand__c", "CurrencyBook__c", "BuyRates__c"),
            MappingRow("Channel__c", "Brand__c", "Brand__c"),
            MappingRow("TaxProfileAssignment__c", "TaxProfile__c", "TaxProfile__c"),
        ]
    )

    positions = {
        tuple(group["members"]): index
        for index, group in enumerate(groups)
    }
    assert positions[("CurrencyBook__c",)] < positions[("Brand__c",)]
    assert positions[("Brand__c",)] < positions[("Channel__c",)]
    assert positions[("TaxProfile__c",)] < positions[("TaxProfileAssignment__c",)]


def test_build_hierarchy_groups_keeps_isolated_objects():
    groups = build_hierarchy_groups(
        [MappingRow("Brand__c", "CurrencyBook__c", "BuyRates__c")],
        object_names=["CurrencyBook__c", "Brand__c", "Ship__c"],
    )

    members = [tuple(group["members"]) for group in groups]
    assert ("Ship__c",) in members


def test_build_hierarchy_analytics_assigns_dependency_waves_and_cycle_phase():
    mappings = [
        MappingRow("B__c", "A__c", "A__c"),
        MappingRow("C__c", "B__c", "B__c"),
        MappingRow("X__c", "Y__c", "Y__c"),
        MappingRow("Y__c", "X__c", "X__c"),
    ]
    groups = build_hierarchy_groups(
        mappings, object_names=["A__c", "B__c", "C__c", "X__c", "Y__c"]
    )
    analytics = build_hierarchy_analytics(
        mappings,
        groups,
        lookup_details={
            "C__c": {
                "record_lookups": ["B__c → B__c"],
                "other_lookups": ["WhatId → A__c, B__c [polymorphic]"],
            }
        },
    )

    assert analytics["A__c"]["dependency_level"] == 1
    assert analytics["A__c"]["is_foundation_object"] == "Yes"
    assert analytics["B__c"]["dependency_level"] == 2
    assert analytics["B__c"]["parent_objects"] == ["A__c"]
    assert analytics["C__c"]["dependency_level"] == 3
    assert analytics["C__c"]["transitive_parent_count"] == 2
    assert analytics["C__c"]["deferred_lookup_count"] == 1
    assert analytics["C__c"]["migration_risk"] == "High"
    assert analytics["X__c"]["migration_phase"] == (
        "Cycle Pass 1 + Deferred Lookup Pass 2"
    )
    assert analytics["X__c"]["deferred_lookup_count"] == 1


def test_generate_object_hierarchy_workbook_writes_expected_sheets():
    connector = _FakeConnector(
        {
            "A__c": {
                "fields": [{
                    "name": "B__c",
                    "type": "reference",
                    "referenceTo": ["B__c"],
                    "cascadeDelete": True,
                    "nillable": False,
                }]
            },
            "B__c": {
                "fields": [{"name": "A__c", "type": "reference", "referenceTo": ["A__c"]}]
            },
        }
    )

    out_path = PROJECT_ROOT / "state" / "test_object_hierarchy.xlsx"
    try:
        result = generate_object_hierarchy_workbook(connector, ["A__c", "B__c"], out_path)

        assert result["mapping_count"] == 2
        assert result["cycle_group_count"] == 1
        assert out_path.exists()

        wb = load_workbook(out_path)
        assert wb.sheetnames == ["ParentIdMappings", "ObjectHierarchy"]
        rows = list(wb["ParentIdMappings"].iter_rows(values_only=True))
        assert rows[0] == (
            "ObjectName",
            "ParentObject",
            "ParentIdField",
            "RelationshipType",
            "IsRequired",
        )
        assert ("A__c", "B__c", "B__c", "Master-Detail", "Yes") in rows[1:]
        assert ("B__c", "A__c", "A__c", "Lookup", "No") in rows[1:]
        hierarchy_rows = list(wb["ObjectHierarchy"].iter_rows(values_only=True))
        assert hierarchy_rows[0] == (
            "Order",
            "GroupOrder",
            "ObjectName",
            "IsCycleGroup",
            "GroupMembers",
            "RecordLookupCount",
            "RecordLookups",
            "OtherLookups",
            "DependencyLevel",
            "ParentObjectCount",
            "ParentObjects",
            "ChildObjectCount",
            "ChildObjects",
            "TransitiveParentCount",
            "CriticalPathDepth",
            "IsFoundationObject",
            "MigrationPhase",
            "ParallelBatch",
            "DeferredLookupCount",
            "MigrationRisk",
        )
        assert hierarchy_rows[1][:5] == (1, 1, "A__c", "Yes", "A__c, B__c")
        assert hierarchy_rows[2][:5] == (2, 1, "B__c", "Yes", "A__c, B__c")
        assert hierarchy_rows[1][5] == 1
        assert hierarchy_rows[2][5] == 1
    finally:
        if out_path.exists():
            try:
                os.remove(out_path)
            except OSError:
                pass


def test_generate_object_hierarchy_skips_global_pseudo_object():
    connector = _ActivityConnector({"Account": {"fields": []}})
    out_path = PROJECT_ROOT / "state" / "test_object_hierarchy_skipped.xlsx"
    try:
        result = generate_object_hierarchy_workbook(
            connector, ["Activity", "Account"], out_path
        )

        assert result["object_count"] == 1
        assert result["skipped_object_count"] == 1
        assert result["skipped_objects"][0]["object"] == "Activity"
        wb = load_workbook(out_path)
        assert wb.sheetnames == [
            "ParentIdMappings",
            "ObjectHierarchy",
            "SkippedObjects",
        ]
    finally:
        if out_path.exists():
            try:
                os.remove(out_path)
            except OSError:
                pass


if __name__ == "__main__":
    import sys as _sys
    from tests._runner import run
    _sys.exit(run(globals()))
