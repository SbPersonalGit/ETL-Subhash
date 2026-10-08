import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.consolidate_salesforce_fk_failures import (
    scan_failed_files,
    write_workbook,
)


def write_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rows)


class ForeignKeyFailureTest(unittest.TestCase):
    def test_scan_recovers_case_and_audits_missing_error_column(self) -> None:
        root = Path.cwd()
        paths: list[Path] = []
        workbook_path: Path | None = None
        try:
            for _ in range(3):
                handle = tempfile.NamedTemporaryFile(
                    prefix="Case_failed_unittest_",
                    suffix=".csv",
                    dir=root,
                    delete=False,
                )
                handle.close()
                paths.append(Path(handle.name))
            write_csv(
                paths[0],
                [
                    ["Id", "AccountId", "OwnerId", "CreatedById", "sf__Error"],
                    [
                        "500xx",
                        "001S900001XN5AnIAL",
                        "005S900000ABCdeIAB",
                        "005S900000ABCdeIAB",
                        "INVALID_FIELD:Foreign key external ID: 001s900001xn5anial "
                        "not found for field External_Id__c in entity Account:--",
                    ],
                    [
                        "500yy",
                        "",
                        "005S900000ABCdeIAB",
                        "005S900000ABCdeIAB",
                        "INVALID_FIELD:Foreign key external ID: 005S900000ABCdeIAB "
                        "not found for field External_Id__c in entity User:--",
                    ],
                ],
            )
            write_csv(paths[1], [["Id", "Name"], ["1", "x"]])
            write_csv(
                paths[2],
                [["Id", "sf__Error"], ["1", "REQUIRED_FIELD_MISSING: Subject"]],
            )

            report = scan_failed_files(root, paths)
            self.assertEqual(len(report.failed_files), 3)
            self.assertEqual(len(report.missing_error_column), 1)
            self.assertEqual(len(report.no_qualifying_errors), 1)
            account = report.records_by_entity["Account"][0]
            self.assertTrue(account.converted_to_lowercase)
            self.assertEqual(account.correct_id, "001S900001XN5AnIAL")
            self.assertEqual(account.found_columns, ["AccountId"])
            user = report.records_by_entity["User"][0]
            self.assertEqual(user.found_columns, ["OwnerId", "CreatedById"])

            workbook_handle = tempfile.NamedTemporaryFile(
                prefix="fk_analysis_unittest_",
                suffix=".xlsx",
                dir=root,
                delete=False,
            )
            workbook_handle.close()
            workbook_path = Path(workbook_handle.name)
            write_workbook(report, workbook_path)
            workbook = load_workbook(workbook_path, read_only=True)
            self.assertIn("Account", workbook.sheetnames)
            self.assertIn("User", workbook.sheetnames)
            self.assertIn("Missing sf__Error", workbook.sheetnames)
            workbook.close()
        finally:
            for path in paths:
                path.unlink(missing_ok=True)
            if workbook_path:
                workbook_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
