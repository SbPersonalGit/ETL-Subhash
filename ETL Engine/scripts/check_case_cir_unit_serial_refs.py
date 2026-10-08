"""
Compare Case.GE_OG_CIR_Unit_Serial_number__c values against GE_Installed_Base__c inventory.

Writes summary + Case Id / serial sheets into the Case archive folder.
"""

from __future__ import annotations

import csv
import json
import time
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

CASE_DIR = Path(
    r"C:\Users\puneetmalhotra\OneDrive - Nagarro\Desktop\migration\BH data\Migrated\Puneet\Templates used for ETL\Case- Archive"
)
IB_CSV = Path(
    r"C:\ETL\ETL_For_Local_Only\state\flat_file_drop\GE_Installed_Base__c\GE_Installed_Base__c.csv"
)
FIELD = "GE_OG_CIR_Unit_Serial_number__c"
OUT_XLSX = CASE_DIR / "GE_OG_CIR_Unit_Serial_number__c_reference_check.xlsx"
OUT_MISSING_CSV = CASE_DIR / "GE_OG_CIR_Unit_Serial_number__c_missing_references.csv"
OUT_ALL_POPULATED_CSV = CASE_DIR / "GE_OG_CIR_Unit_Serial_number__c_case_id_serial_populated.csv"
OUT_SUMMARY_JSON = CASE_DIR / "GE_OG_CIR_Unit_Serial_number__c_reference_check.summary.json"


def case_files(folder: Path) -> list[Path]:
    files = []
    for path in sorted(folder.glob("*.csv")):
        name = path.name.lower()
        if name.startswith("unique case"):
            continue
        if "missing" in name or "invalid" in name:
            continue
        if FIELD.lower().replace("__c", "") in name.lower():
            continue
        if "reference_check" in name or "case_id_serial" in name:
            continue
        files.append(path)
    return files


def load_ib_ids(path: Path) -> set[str]:
    ids: set[str] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or "Id" not in reader.fieldnames:
            raise SystemExit(f"GE_Installed_Base inventory missing Id column: {path}")
        for row in reader:
            value = (row.get("Id") or "").strip()
            if value:
                ids.add(value)
                # also accept 15-char form if 18 provided
                if len(value) == 18:
                    ids.add(value[:15])
    return ids


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    started = time.time()
    files = case_files(CASE_DIR)
    if not files:
        raise SystemExit(f"No Case CSV files found in {CASE_DIR}")
    if not IB_CSV.exists():
        raise SystemExit(f"Missing inventory: {IB_CSV}")

    print(f"Loading IB inventory: {IB_CSV}")
    ib_ids = load_ib_ids(IB_CSV)
    print(f"IB unique Ids: {len(ib_ids):,}")
    print(f"Case files: {len(files)}")

    total_rows = 0
    populated_rows = 0
    blank_rows = 0
    found_rows = 0
    missing_rows = 0
    distinct_values: set[str] = set()
    distinct_found: set[str] = set()
    distinct_missing: set[str] = set()
    per_file = []

    missing_pairs: list[tuple[str, str, str]] = []  # case_id, serial, source_file
    # Stream all populated pairs to CSV to avoid huge memory
    OUT_ALL_POPULATED_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_MISSING_CSV.open("w", encoding="utf-8", newline="") as missing_handle, OUT_ALL_POPULATED_CSV.open(
        "w", encoding="utf-8", newline=""
    ) as all_handle:
        missing_writer = csv.writer(missing_handle)
        all_writer = csv.writer(all_handle)
        missing_writer.writerow(["Case_Id", FIELD, "SourceFile"])
        all_writer.writerow(["Case_Id", FIELD, "MatchStatus", "SourceFile"])

        for path in files:
            print(f"Scanning {path.name}...")
            file_rows = file_populated = file_missing = file_found = 0
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise SystemExit(f"No headers in {path}")
                if "Id" not in reader.fieldnames:
                    raise SystemExit(f"Id column missing in {path}")
                if FIELD not in reader.fieldnames:
                    raise SystemExit(f"{FIELD} column missing in {path}")
                for row in reader:
                    file_rows += 1
                    total_rows += 1
                    case_id = (row.get("Id") or "").strip()
                    serial = (row.get(FIELD) or "").strip()
                    if not serial:
                        blank_rows += 1
                        continue
                    populated_rows += 1
                    file_populated += 1
                    distinct_values.add(serial)
                    # normalize 18->15 for membership if needed
                    matched = serial in ib_ids or (len(serial) == 18 and serial[:15] in ib_ids)
                    if matched:
                        found_rows += 1
                        file_found += 1
                        distinct_found.add(serial)
                        status = "FOUND"
                    else:
                        missing_rows += 1
                        file_missing += 1
                        distinct_missing.add(serial)
                        status = "MISSING"
                        missing_pairs.append((case_id, serial, path.name))
                        missing_writer.writerow([case_id, serial, path.name])
                    all_writer.writerow([case_id, serial, status, path.name])
            per_file.append(
                {
                    "file": path.name,
                    "rows": file_rows,
                    "populated": file_populated,
                    "found": file_found,
                    "missing": file_missing,
                }
            )
            print(
                f"  rows={file_rows:,} populated={file_populated:,} "
                f"found={file_found:,} missing={file_missing:,}"
            )

    # Excel workbook with summary + missing + distinct missing
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    for metric, value in [
        ("CaseFolder", str(CASE_DIR)),
        ("InstalledBaseInventory", str(IB_CSV)),
        ("LookupField", FIELD),
        ("ReferencedObject", "GE_Installed_Base__c"),
        ("MatchKey", "GE_Installed_Base__c.Id"),
        ("CaseFilesScanned", len(files)),
        ("CaseRowsTotal", total_rows),
        ("LookupBlankRows", blank_rows),
        ("LookupPopulatedRows", populated_rows),
        ("LookupFoundRows", found_rows),
        ("LookupMissingRows", missing_rows),
        ("DistinctLookupValues", len(distinct_values)),
        ("DistinctFoundValues", len(distinct_found)),
        ("DistinctMissingValues", len(distinct_missing)),
        ("InstalledBaseUniqueIds", len(ib_ids)),
        ("StillMissingReferences", "YES" if missing_rows > 0 else "NO"),
        ("ElapsedSeconds", round(time.time() - started, 1)),
    ]:
        summary.append([metric, value])
    style_header(summary)

    per_file_ws = wb.create_sheet("PerFile")
    per_file_ws.append(["SourceFile", "Rows", "Populated", "Found", "Missing"])
    for item in per_file:
        per_file_ws.append(
            [item["file"], item["rows"], item["populated"], item["found"], item["missing"]]
        )
    style_header(per_file_ws)

    missing_ws = wb.create_sheet("Missing_CaseId_Serial")
    missing_ws.append(["Case_Id", FIELD, "SourceFile"])
    # Excel row limit safety
    max_excel_rows = 1_000_000
    for idx, (case_id, serial, source) in enumerate(missing_pairs):
        if idx >= max_excel_rows - 1:
            missing_ws.append(["TRUNCATED", f"See CSV {OUT_MISSING_CSV.name}", ""])
            break
        missing_ws.append([case_id, serial, source])
    style_header(missing_ws)

    distinct_ws = wb.create_sheet("Distinct_Missing_Serials")
    distinct_ws.append([FIELD, "CaseRowCount"])
    missing_value_counts = Counter(serial for _, serial, _ in missing_pairs)
    for serial, count in missing_value_counts.most_common():
        distinct_ws.append([serial, count])
    style_header(distinct_ws)

    # Also a sheet with Case Id + serial for ALL populated is too large for Excel;
    # point to CSV instead.
    note_ws = wb.create_sheet("Notes")
    note_ws.append(["Note"])
    note_ws.append(
        [
            f"Full Case Id + {FIELD} for every populated row (FOUND and MISSING) is in: {OUT_ALL_POPULATED_CSV.name}"
        ]
    )
    note_ws.append([f"Missing-only Case Id + {FIELD} CSV: {OUT_MISSING_CSV.name}"])
    note_ws.append(["Inventory match used GE_Installed_Base__c.Id (file has no Legacy_ID__c column)."])
    style_header(note_ws)

    wb.save(OUT_XLSX)

    summary_payload = {
        "case_folder": str(CASE_DIR),
        "installed_base_inventory": str(IB_CSV),
        "lookup_field": FIELD,
        "referenced_object": "GE_Installed_Base__c",
        "match_key": "GE_Installed_Base__c.Id",
        "case_files_scanned": len(files),
        "case_rows_total": total_rows,
        "lookup_blank_rows": blank_rows,
        "lookup_populated_rows": populated_rows,
        "lookup_found_rows": found_rows,
        "lookup_missing_rows": missing_rows,
        "distinct_lookup_values": len(distinct_values),
        "distinct_found_values": len(distinct_found),
        "distinct_missing_values": len(distinct_missing),
        "installed_base_unique_ids": len(ib_ids),
        "still_missing_references": missing_rows > 0,
        "outputs": {
            "xlsx": str(OUT_XLSX),
            "missing_csv": str(OUT_MISSING_CSV),
            "all_populated_csv": str(OUT_ALL_POPULATED_CSV),
        },
        "elapsed_seconds": round(time.time() - started, 1),
    }
    OUT_SUMMARY_JSON.write_text(json.dumps(summary_payload, indent=2), encoding="utf-8")

    print("\n=== RESULT ===")
    print(json.dumps(summary_payload, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
