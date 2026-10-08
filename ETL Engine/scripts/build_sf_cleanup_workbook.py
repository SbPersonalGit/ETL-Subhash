"""One-off: build Salesforce_Data_Cleanup_Framework.xlsx in project root."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Salesforce_Data_Cleanup_Framework.xlsx"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border


def autosize(ws, max_width: int = 55) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        maxlen = min(max(len(str(c.value or "")) for c in col), max_width)
        ws.column_dimensions[letter].width = max(maxlen + 2, 12)


def main() -> None:
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Overview"
    ws0["A1"] = "Salesforce Account & Contact cleanup — framework workbook"
    ws0["A1"].font = Font(bold=True, size=14)
    ws0["A3"] = "How to use & pre-cleanup backup"
    ws0["A3"].font = Font(bold=True)
    overview_end_row = 24
    ws0.merge_cells(f"A4:F{overview_end_row}")
    ws0["A4"] = (
        "HOW TO USE\n"
        "Use each sheet as a working template. Fill status/notes as you execute.\n"
        "• Challenges — problem inventory to align stakeholders.\n"
        "• Solution phases — end-to-end delivery model.\n"
        "• Discovery checklist — data & integration baseline.\n"
        "• Downstream systems — merge/delete blast radius.\n"
        "• Matching rules — duplicate logic & survivorship (includes different owners on duplicates).\n"
        "• Delivery waves — scoped rollout with risk & timing.\n"
        "• KPIs — before/after measurement.\n"
        "\n"
        "BEFORE CLEANUP — BACKUP & RECOVERABILITY\n"
        "Treat backup (or equivalent restore capability) as mandatory before merges, bulk updates, or deletes.\n"
        "• Export or back up objects you will touch, including Salesforce Ids, with a clear timestamp / wave id.\n"
        "• If the org uses a backup product, confirm record-level restore and retention vs. one-off CSV exports.\n"
        "• Rehearse waves in sandbox first; rehearsal reduces risk but does not replace production backup.\n"
        "• Merges retire duplicate Ids — plan Id mapping (old → master) and downstream key updates; rollback is rarely a single re-import.\n"
        "• Deletes/archives: add approvals; prefer soft-archive (status, stop sync) where hard delete is risky or policy-restricted.\n"
        "• Document automation/integrations that could recreate bad data after a restore."
    )
    ws0["A4"].alignment = wrap
    for r in range(4, overview_end_row + 1):
        for c in range(1, 7):
            ws0.cell(r, c).border = border

    ws1 = wb.create_sheet("Challenges")
    ws1.append(["ID", "Challenge", "Why it matters", "Typical symptoms / signals"])
    for r in [
        (
            "C01",
            "Duplicates & near-duplicates",
            "Wrong revenue rollups, bad routing, wasted outreach.",
            "High Account/Contact counts; same domain/legal name variants.",
        ),
        (
            "C02",
            "Unclear Account/Contact grain",
            "Merges applied to wrong entity type; rework and mistrust.",
            "Mixed use of HQ vs site vs named account without policy.",
        ),
        (
            "C03",
            "SF merge/delete technical limits",
            "Failed batches, automation errors, partial fixes.",
            "Merge failures; related lists blocked; required fields missing.",
        ),
        (
            "C04",
            "Downstream system impact",
            "Broken keys in ERP/warehouse; duplicate billing or support cases.",
            "Integration errors after merges; orphaned external IDs.",
        ),
        (
            "C05",
            "Automation & integrations recreate bad data",
            "Cleanup gains erode within days/weeks.",
            "Spike in duplicates post-integration runs.",
        ),
        (
            "C06",
            "Ownership & territory sensitivity",
            "Stalled decisions; revenue disputes.",
            "Pushback on reassignment or survivor record choice.",
        ),
        (
            "C07",
            "Low data completeness / legacy imports",
            "Noisy matching; false merge risk.",
            "Sparse addresses; junk picklists; stale sources.",
        ),
        (
            "C08",
            "Global scale & prioritization",
            "QA overload; inconsistent regional quality.",
            "No scoped waves; everything is priority one.",
        ),
        (
            "C09",
            "Audit & compliance",
            "Inability to explain merges/deletes; regulatory exposure.",
            "Missing approval trail; GDPR/retention ambiguity.",
        ),
        (
            "C10",
            "External golden data limitations",
            "Wrong merges when vendor data does not fit all markets.",
            "MDM/D&B mismatches vs local entities.",
        ),
    ]:
        ws1.append(list(r))
    style_header(ws1)
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws1)

    ws2 = wb.create_sheet("Solution phases")
    ws2.append(["Phase", "Activities (examples)", "Key outputs", "Suggested owners"])
    for r in [
        (
            "1. Governance & definitions",
            "Agree Account/Contact grain; merge/archive/delete policy; RACI.",
            "Signed definitions; risk appetite; steering cadence.",
            "Sales Ops, Legal/Privacy, IT, Data",
        ),
        (
            "2. Discovery & baseline",
            "Object inventory; integration map; sample profiling; KPI baseline.",
            "System map; data quality snapshot; prioritized problem statements.",
            "Data analyst, Integration leads, SF admin",
        ),
        (
            "3. Matching & survivorship design",
            "Tiered match keys; confidence bands; field-level survive rules.",
            "Match spec; exception criteria; QA sampling plan.",
            "Data analyst, MDM owner, SF admin",
        ),
        (
            "4. Pilot wave",
            "Narrow scope; dry-run; reconcile with downstream replay/sandbox.",
            "Go/no-go; refined rules; runbook updates.",
            "Delivery pod + business sponsor",
        ),
        (
            "5. Scaled execution",
            "Waves by region/segment/source; exports & id maps; comms.",
            "Reduced duplicates; stable integrations; audit trail.",
            "Ops + IT + stewards",
        ),
        (
            "6. Prevention & monitoring",
            "Ingress validation; duplicate prevention; dashboards; SLAs.",
            "Sustained KPIs; fewer reactive fire drills.",
            "Sales Ops + IT + COE",
        ),
    ]:
        ws2.append(list(r))
    style_header(ws2)
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws2)

    ws3 = wb.create_sheet("Discovery checklist")
    ws3.append(
        ["Area", "Checklist item", "Status (Pending/In progress/Done)", "Notes / evidence link", "Owner"]
    )
    for it in [
        ("Salesforce model", "Record types, layouts, required fields for Account/Contact"),
        ("Salesforce model", "Automation: flows, PB, Apex, DLRS/rollups affecting merge"),
        ("Salesforce model", "Duplicate Jobs / matching rules configuration"),
        ("Data sources", "List of systems creating/updating Accounts & Contacts"),
        ("Data sources", "Historical migration batches and known bad sources"),
        ("Identifiers", "Canonical keys: ERP Id, MDM Id, DUNS, domain, national Id, etc."),
        ("Identifiers", "Idempotency: how upserts are keyed in each integration"),
        ("Integrations", "Error logs / replay behavior on update vs merge scenarios"),
        ("Analytics", "Warehouse facts keyed by SF Id vs business natural keys"),
        ("Legal / privacy", "Retention, deletion, and consent rules by region"),
        (
            "Governance / Sales Ops",
            "Written policy when duplicate pair has different Owner (survivorship; RevOps/comp; exceptions queue)",
        ),
        (
            "Governance / Sales Ops",
            "How open Opportunities/Contacts will be realigned if Account owner changes on merge",
        ),
        (
            "Salesforce model",
            "Account Team / overlay / sharing plan so second rep stays visible when Owner must be single-valued",
        ),
    ]:
        ws3.append([it[0], it[1], "", "", ""])
    style_header(ws3)
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws3)

    ws4 = wb.create_sheet("Downstream systems")
    ws4.append(
        [
            "System",
            "Business owner",
            "Primary match key(s)",
            "Impact if Account merged",
            "Impact if record deleted/archived",
            "Mitigation / notes",
        ]
    )
    ws4.append(
        [
            "Example: ERP",
            "",
            "Account.CustomerNumber",
            "May duplicate customer or break billing link",
            "Open invoices / contracts may block delete",
            "Map surviving Id; pre-flight validation",
        ]
    )
    style_header(ws4)
    for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws4)

    ws5 = wb.create_sheet("Matching rules")
    ws5.append(
        [
            "Object",
            "Rule name / tier",
            "Match logic (fields & fuzzy)",
            "Confidence (High/Med/Low)",
            "Survivorship (which value wins)",
            "Exception / manual review trigger",
        ]
    )
    ws5.append(
        [
            "Account",
            "Tier A — deterministic",
            "Exact External_Id or ERP Id match",
            "High",
            "Longest tenure; or Finance-owned Id as master",
            "Conflict on country/segment — steward",
        ]
    )
    ws5.append(
        [
            "Account",
            "Tier B — probabilistic",
            "Legal name + postal + country (fuzzy)",
            "Medium",
            "Populate blanks from best secondary",
            "ACV > threshold — manual",
        ]
    )
    ws5.append(
        [
            "Contact",
            "Tier A",
            "Email exact (normalized)",
            "High",
            "Keep complete profile + consent flags",
            "Multiple Accounts same person — policy",
        ]
    )
    ws5.append(
        [
            "Account",
            "Owner — different owners on duplicate pair",
            "Dup pair identified; Account.OwnerId differs between master candidate rows",
            "High",
            "Default: master record’s owner. Common tie-breaks: largest open pipe; territory model; latest activity; steward.",
            "Both have material open pipe; strategic account; comp dispute — manual queue + comms",
        ]
    )
    ws5.append(
        [
            "Contact",
            "Owner — different owners on duplicate pair",
            "Dup pair identified; Contact.OwnerId differs (or inherited vs explicit)",
            "High",
            "Align with Account-owner policy; keep richer engagement history on survivor; reassign open Contact tasks.",
            "Partner/contact ownership splits; GDPR-sensitive contact — manual",
        ]
    )
    ws5.append(
        [
            "Account",
            "Mitigation — single Owner field",
            "N/A (post-merge)",
            "Policy",
            "Add non-owner to Account Team / overlay; realign open Opps to agreed owner before/after merge.",
            "Document in runbook; train reps before wave",
        ]
    )
    style_header(ws5)
    for row in ws5.iter_rows(min_row=2, max_row=ws5.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws5)

    ws6 = wb.create_sheet("Delivery waves")
    ws6.append(
        [
            "Wave ID",
            "Scope (region/segment/source)",
            "Est. record volume",
            "Risk rating (L/M/H)",
            "Planned start",
            "Status",
            "Exit criteria / notes",
        ]
    )
    ws6.append(
        [
            "W01",
            "e.g. Legacy CRM import — NA enterprise",
            "",
            "High",
            "",
            "Planned",
            "Pilot: sandbox replay + 1% prod sample",
        ]
    )
    style_header(ws6)
    for row in ws6.iter_rows(min_row=2, max_row=ws6.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws6)

    ws7 = wb.create_sheet("KPIs")
    ws7.append(["Metric", "Baseline (current)", "Target", "How measured (tool/report)", "Refresh cadence"])
    for m in [
        ("Duplicate Account clusters (est.)", "", "", "Dedupe report / graph match", "Weekly"),
        ("% Accounts with trusted external key", "", "", "Profile on key field population", "Monthly"),
        ("Integration errors tied to customer master", "", "", "Log analysis tagged by object", "Weekly"),
        ("Time to resolve steward queue (median)", "", "", "Ticketing / task data", "Weekly"),
        ("Records touched per wave (merge/update/archive)", "", "", "Wave execution log", "Per wave"),
    ]:
        ws7.append(list(m))
    style_header(ws7)
    for row in ws7.iter_rows(min_row=2, max_row=ws7.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = wrap
    autosize(ws7)

    try:
        wb.save(OUT)
        print(f"Wrote {OUT}")
    except PermissionError:
        alt = OUT.with_name(f"{OUT.stem}_regenerated{OUT.suffix}")
        wb.save(alt)
        print(f"Original open or locked; wrote {alt}")


if __name__ == "__main__":
    main()
