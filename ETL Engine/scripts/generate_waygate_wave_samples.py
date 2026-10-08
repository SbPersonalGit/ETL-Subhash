"""Generate target-aware sample CSVs for the first objects in a hierarchy wave."""

from __future__ import annotations

import argparse
import csv
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from core.connector import connector_from_env
from scripts.generate_waygate_standard_samples import (
    is_writable,
    scalar_value,
    stable_source_id,
    write_csv,
)


def wave_objects(workbook: Path, wave: str, count: int) -> list[dict]:
    sheet = load_workbook(workbook, read_only=True, data_only=True)["ObjectHierarchy"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    matches = [
        dict(zip(headers, values))
        for values in rows
        if str(dict(zip(headers, values)).get("ParallelBatch") or "").strip() == wave
    ]
    matches.sort(key=lambda item: int(item["Order"]))
    if len(matches) < count:
        raise ValueError(f"{wave} contains only {len(matches)} objects; requested {count}")
    return matches[:count]


def select_fields(
    fields: list[dict], coverage: float, minimum_fields: int, include_owner: bool
) -> tuple[list[dict], int, bool, int]:
    writable = [field for field in fields if is_writable(field)]
    by_name = {str(field["name"]): field for field in writable}
    legacy_from_id = "Legacy_ID__c" in by_name
    required = [
        field
        for field in writable
        if not field.get("nillable", True) and not field.get("defaultedOnCreate", False)
    ]
    required_refs = [field["name"] for field in required if field.get("type") == "reference"]
    if required_refs:
        raise ValueError("required parent data is not in this wave: " + ", ".join(required_refs))
    required_names = {field["name"] for field in required}
    candidates = [
        field
        for field in writable
        if field["name"] not in required_names
        and field["name"] != "Legacy_ID__c"
        and field.get("type") != "reference"
    ]
    if include_owner and "OwnerId" in by_name:
        candidates.insert(0, by_name["OwnerId"])
    requested = max(math.ceil(len(writable) * coverage), minimum_fields)
    target = requested - int(legacy_from_id)
    selected = [*required, *candidates[: max(0, target - len(required))]]
    # A dependent picklist is not safely loadable without its controlling
    # field. Prune repeatedly because dependent picklists may form a chain.
    while True:
        selected_names = {str(field["name"]) for field in selected}
        unsafe = {
            str(field["name"])
            for field in selected
            if field.get("dependentPicklist")
            and str(field.get("controllerName") or "") not in selected_names
        }
        if not unsafe:
            break
        selected = [field for field in selected if str(field["name"]) not in unsafe]
    return selected, len(writable), legacy_from_id, requested


def dependent_picklist_value(
    field: dict,
    row: dict[str, str],
    picklist_metadata: dict[str, dict],
    row_index: int,
) -> str:
    config = picklist_metadata.get(str(field["name"])) or {}
    controller = str(field.get("controllerName") or "")
    controller_value = str(row.get(controller) or "")
    controller_index = (config.get("controller_values") or {}).get(controller_value)
    valid = [
        value
        for value in sorted(config.get("values") or [])
        if controller_index in (config.get("valid_for") or {}).get(value, set())
    ]
    if not valid:
        raise ValueError(
            f"{field['name']}: no target picklist value is valid for "
            f"{controller}={controller_value!r}"
        )
    return valid[row_index % len(valid)]


def dependency_order(fields: list[dict]) -> list[dict]:
    """Order fields so each dependent picklist follows its controller."""
    pending = list(fields)
    ordered: list[dict] = []
    emitted: set[str] = set()
    while pending:
        ready = [
            field
            for field in pending
            if not field.get("dependentPicklist")
            or str(field.get("controllerName") or "") in emitted
        ]
        if not ready:
            names = ", ".join(str(field["name"]) for field in pending)
            raise ValueError(f"Unresolvable dependent-picklist order: {names}")
        for field in ready:
            ordered.append(field)
            emitted.add(str(field["name"]))
            pending.remove(field)
    return ordered


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--profile", default="default")
    parser.add_argument("--records-per-object", type=int, default=100)
    parser.add_argument("--coverage", type=float, default=0.20)
    parser.add_argument("--minimum-fields", type=int, default=15)
    parser.add_argument("--exclude-owner", action="store_true")
    parser.add_argument("--wave", default="Wave 1")
    parser.add_argument("--object-count", type=int, default=10)
    parser.add_argument(
        "--hierarchy-workbook",
        type=Path,
        default=Path("temp/hierarchy_edit/waygate_object_hierarchy.xlsx"),
    )
    args = parser.parse_args()
    if not 0 < args.coverage <= 1:
        raise ValueError("--coverage must be greater than 0 and no more than 1")
    if args.records_per_object < 1 or args.object_count < 1 or args.minimum_fields < 1:
        raise ValueError("record and object counts must be positive")

    hierarchy_rows = wave_objects(args.hierarchy_workbook, args.wave, args.object_count)
    object_names = [str(row["ObjectName"]) for row in hierarchy_rows]
    output = args.output.resolve()
    upload_dir = output / "upload_files"
    upload_dir.mkdir(parents=True, exist_ok=True)
    target = connector_from_env("TARGET", args.profile).connect()
    owner = target.query(
        "SELECT Id, Username FROM User WHERE IsActive = true AND UserType = 'Standard' "
        "ORDER BY LastLoginDate DESC NULLS LAST LIMIT 1"
    )
    if not owner and not args.exclude_owner:
        raise RuntimeError("No active target User is available for OwnerId mapping")
    owner_record = owner[0] if owner else None
    manifest = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "hierarchy_source": str(args.hierarchy_workbook.resolve()),
        "wave": args.wave,
        "requested_coverage_percent": args.coverage * 100,
        "records_per_object": args.records_per_object,
        "total_records": args.records_per_object * len(object_names),
        "load_order": object_names,
        "target_org_owner": owner_record if not args.exclude_owner else None,
        "objects": [],
    }

    for hierarchy, object_name in zip(hierarchy_rows, object_names):
        describe = target.describe_object(object_name)
        if not describe.get("createable"):
            raise PermissionError(f"Target profile cannot create {object_name}")
        selected, writable_count, legacy_from_id, requested_fields = select_fields(
            describe.get("fields") or [],
            args.coverage,
            args.minimum_fields,
            not args.exclude_owner,
        )
        default_record_type_id = next(
            (
                str(info.get("recordTypeId") or "")
                for info in (describe.get("recordTypeInfos") or [])
                if info.get("defaultRecordTypeMapping") and info.get("available", True)
            ),
            "",
        )
        picklist_metadata = (
            target.describe_record_type_picklists(object_name, default_record_type_id)
            if default_record_type_id
            else {}
        )
        rows = []
        for row_index in range(args.records_per_object):
            row = {"Id": stable_source_id(object_name, row_index)}
            ordered_fields = dependency_order(selected)
            for field in ordered_fields:
                if field["name"] == "OwnerId":
                    row["OwnerId"] = str(owner_record["Id"])
                elif field.get("dependentPicklist"):
                    row[str(field["name"])] = dependent_picklist_value(
                        field, row, picklist_metadata, row_index
                    )
                else:
                    row[str(field["name"])] = scalar_value(object_name, field, row_index)
            rows.append(row)
        headers = ["Id", *[str(field["name"]) for field in selected]]
        write_csv(upload_dir / f"{object_name}.csv", headers, rows)
        populated = sum(any(str(row.get(name) or "").strip() for row in rows) for name in headers[1:])
        covered = populated + int(legacy_from_id)
        actual = covered / writable_count if writable_count else 0
        if actual < args.coverage:
            raise RuntimeError(f"{object_name} coverage {actual:.1%} is below target")
        manifest["objects"].append({
            "order": int(hierarchy["Order"]),
            "object": object_name,
            "csv": f"upload_files/{object_name}.csv",
            "records": len(rows),
            "writable_non_calculated_fields": writable_count,
            "selected_fields": len(selected) + int(legacy_from_id),
            "populated_fields": covered,
            "coverage_percent": round(actual * 100, 1),
            "requested_minimum_fields": requested_fields,
            "eligible_field_shortfall": max(0, requested_fields - covered),
            "record_lookup_count": int(hierarchy.get("RecordLookupCount") or 0),
            "global_target_references": ["OwnerId"] if any(field["name"] == "OwnerId" for field in selected) else [],
        })

    (output / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    with (output / "coverage.csv").open("w", encoding="utf-8", newline="") as handle:
        columns = ["order", "object", "records", "writable_non_calculated_fields", "selected_fields", "populated_fields", "coverage_percent", "requested_minimum_fields", "eligible_field_shortfall", "record_lookup_count"]
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows({key: item[key] for key in columns} for item in manifest["objects"])
    readme = [
        f"# Waygate {args.wave} sample package",
        "",
        f"Contains {len(object_names)} objects, {args.records_per_object} records per object, and {manifest['total_records']} total records.",
        f"Selection target: at least {args.coverage:.0%} or {args.minimum_fields} fields, whichever is higher.",
        "When an object has fewer eligible loadable fields, every eligible field is populated and the shortfall is recorded in coverage.csv.",
        "OwnerId values reference the active target-org User recorded in manifest.json.",
        "These foundation objects have no record lookup dependencies in the hierarchy workbook.",
        "",
        "Load order:",
        "",
        *[f"{index}. `upload_files/{name}.csv`" for index, name in enumerate(object_names, 1)],
    ]
    (output / "README.md").write_text("\n".join(readme) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
