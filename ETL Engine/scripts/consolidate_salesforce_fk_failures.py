"""Collect failed Salesforce CSVs and analyze foreign-key External ID errors.

The source tree is scanned recursively. Failed CSVs are copied while preserving
their relative paths, and an Excel workbook groups parsed errors by referenced
Salesforce object (Account, User, and so on).
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ERROR_PREFIX = "INVALID_FIELD:Foreign key external ID"
ERROR_RE = re.compile(
    r"^INVALID_FIELD:Foreign key external ID:\s*"
    r"(?P<record_id>[A-Za-z0-9]{15,18})\s+not found for field\s+"
    r"(?P<external_field>\S+)\s+in entity\s+"
    r"(?P<entity>[A-Za-z_][A-Za-z0-9_]*)",
    re.IGNORECASE,
)
SF_ID_RE = re.compile(r"^[A-Za-z0-9]{15}(?:[A-Za-z0-9]{3})?$")


@dataclass
class ErrorRecord:
    entity: str
    found_id: str
    converted_to_lowercase: bool
    correct_id: str
    found_columns: list[str]
    values_by_column: dict[str, str]
    correct_id_status: str
    source_file: str
    source_row: int
    external_id_field: str
    sf_error: str


@dataclass
class ScanReport:
    source_root: Path
    failed_files: list[Path]
    total_rows: int = 0
    rows_with_error: int = 0
    qualifying_errors: int = 0
    records_by_entity: dict[str, list[ErrorRecord]] = field(
        default_factory=lambda: defaultdict(list)
    )
    missing_error_column: list[Path] = field(default_factory=list)
    no_qualifying_errors: list[Path] = field(default_factory=list)
    unparsed_errors: list[tuple[Path, int, str]] = field(default_factory=list)


def is_all_lowercase_id(value: str) -> bool:
    return any(char.isalpha() for char in value) and value == value.lower()


def choose_correct_id(found_id: str, row: dict[str, str]) -> tuple[str, list[str], str]:
    matches: list[tuple[str, str]] = []
    for column, raw_value in row.items():
        value = str(raw_value or "").strip()
        if SF_ID_RE.fullmatch(value) and value.casefold() == found_id.casefold():
            matches.append((column, value))
    columns = [column for column, _ in matches]
    candidates = [value for _, value in matches]
    mixed_case = [value for value in candidates if not is_all_lowercase_id(value)]
    if mixed_case:
        counts = Counter(mixed_case)
        correct_id, _ = counts.most_common(1)[0]
        status = "recovered_from_row" if correct_id != found_id else "already_correct"
        if len({value for value in mixed_case}) > 1:
            status = "ambiguous_case_variants"
        return correct_id, columns, status
    if not is_all_lowercase_id(found_id):
        return found_id, columns, "already_correct"
    return "", columns, "correct_case_not_found"


def find_failed_csvs(source_root: Path) -> list[Path]:
    return sorted(
        (
            path
            for path in source_root.rglob("*.csv")
            if "failed" in path.name.casefold()
        ),
        key=lambda path: str(path).casefold(),
    )


def scan_failed_files(source_root: Path, failed_files: list[Path]) -> ScanReport:
    report = ScanReport(source_root=source_root, failed_files=failed_files)
    for index, path in enumerate(failed_files, start=1):
        relative = path.relative_to(source_root)
        print(f"[{index}/{len(failed_files)}] Reading {relative}")
        file_qualifying = 0
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            error_column = next(
                (header for header in headers if header.casefold() == "sf__error"), None
            )
            if not error_column:
                report.missing_error_column.append(path)
                continue
            for row_number, row in enumerate(reader, start=2):
                report.total_rows += 1
                error = str(row.get(error_column) or "").strip()
                if error:
                    report.rows_with_error += 1
                if not error.casefold().startswith(ERROR_PREFIX.casefold()):
                    continue
                report.qualifying_errors += 1
                file_qualifying += 1
                match = ERROR_RE.match(error)
                if not match:
                    report.unparsed_errors.append((path, row_number, error))
                    continue
                found_id = match.group("record_id")
                correct_id, found_columns, status = choose_correct_id(found_id, row)
                entity = match.group("entity")
                report.records_by_entity[entity].append(
                    ErrorRecord(
                        entity=entity,
                        found_id=found_id,
                        converted_to_lowercase=is_all_lowercase_id(found_id),
                        correct_id=correct_id,
                        found_columns=found_columns,
                        values_by_column={
                            column: str(row.get(column) or "").strip()
                            for column in found_columns
                        },
                        correct_id_status=status,
                        source_file=str(relative),
                        source_row=row_number,
                        external_id_field=match.group("external_field"),
                        sf_error=error,
                    )
                )
        if error_column and file_qualifying == 0:
            report.no_qualifying_errors.append(path)
    return report


def safe_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name).strip()[:31] or "Object"
    candidate = base
    counter = 2
    while candidate.casefold() in used:
        suffix = f"_{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def append_header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_workbook(report: ScanReport, output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    append_header(summary, ["Metric", "Count"])
    unresolved = sum(
        record.correct_id_status == "correct_case_not_found"
        for records in report.records_by_entity.values()
        for record in records
    )
    recovered = sum(
        record.correct_id_status == "recovered_from_row"
        for records in report.records_by_entity.values()
        for record in records
    )
    for metric, count in [
        ("Failed CSV files", len(report.failed_files)),
        ("Failed files missing sf__Error", len(report.missing_error_column)),
        ("Failed files with no qualifying FK error", len(report.no_qualifying_errors)),
        ("Rows scanned", report.total_rows),
        ("Rows with sf__Error", report.rows_with_error),
        ("Qualifying foreign-key errors", report.qualifying_errors),
        ("Unparsed qualifying errors", len(report.unparsed_errors)),
        ("Referenced objects", len(report.records_by_entity)),
        ("Lowercase IDs recovered", recovered),
        ("Correct-case IDs not found", unresolved),
    ]:
        summary.append([metric, count])

    used = {"summary"}
    for entity in sorted(report.records_by_entity, key=str.casefold):
        records = report.records_by_entity[entity]
        dynamic_columns = list(
            dict.fromkeys(column for record in records for column in record.found_columns)
        )
        sheet = workbook.create_sheet(safe_sheet_name(entity, used))
        headers = [
            "FoundRecordId",
            "ConvertedToLowercase",
            "CorrectRecordId",
            "CorrectValueFoundInColumns",
            "CorrectIdStatus",
            "SourceFile",
            "SourceRow",
            "ExternalIdField",
            "sf__Error",
            *dynamic_columns,
        ]
        append_header(sheet, headers)
        for record in records:
            sheet.append(
                [
                    record.found_id,
                    "Yes" if record.converted_to_lowercase else "No",
                    record.correct_id,
                    ", ".join(record.found_columns),
                    record.correct_id_status,
                    record.source_file,
                    record.source_row,
                    record.external_id_field,
                    record.sf_error,
                    *(record.values_by_column.get(column, "") for column in dynamic_columns),
                ]
            )

    missing = workbook.create_sheet(safe_sheet_name("Missing sf__Error", used))
    append_header(missing, ["RelativeFile", "FullPath"])
    for path in report.missing_error_column:
        missing.append([str(path.relative_to(report.source_root)), str(path)])

    no_match = workbook.create_sheet(safe_sheet_name("No qualifying errors", used))
    append_header(no_match, ["RelativeFile", "FullPath"])
    for path in report.no_qualifying_errors:
        no_match.append([str(path.relative_to(report.source_root)), str(path)])

    unparsed = workbook.create_sheet(safe_sheet_name("Unparsed errors", used))
    append_header(unparsed, ["RelativeFile", "SourceRow", "sf__Error"])
    for path, row_number, error in report.unparsed_errors:
        unparsed.append([str(path.relative_to(report.source_root)), row_number, error])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def copy_failed_files(report: ScanReport, destination: Path) -> None:
    for path in report.failed_files:
        target = destination / path.relative_to(report.source_root)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)


def write_missing_error_report(report: ScanReport, output: Path) -> None:
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["RelativeFile", "FullPath"])
        for path in report.missing_error_column:
            writer.writerow([str(path.relative_to(report.source_root)), str(path)])


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_root", type=Path, help="Root folder to scan recursively")
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Default: sibling '<source folder> - failed analysis' directory",
    )
    parser.add_argument("--no-copy", action="store_true", help="Do not copy failed CSV files")
    args = parser.parse_args()
    args.source_root = args.source_root.expanduser().resolve()
    if not args.source_root.is_dir():
        parser.error(f"Source folder was not found: {args.source_root}")
    if args.output_dir is None:
        args.output_dir = args.source_root.with_name(
            f"{args.source_root.name} - failed analysis"
        )
    else:
        args.output_dir = args.output_dir.expanduser().resolve()
    return args


def main() -> int:
    args = parse_args()
    failed_files = find_failed_csvs(args.source_root)
    if not failed_files:
        raise SystemExit(f"No failed CSV files found below: {args.source_root}")
    report = scan_failed_files(args.source_root, failed_files)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    if not args.no_copy:
        copy_failed_files(report, args.output_dir / "failed_files")
    workbook = args.output_dir / "foreign_key_failure_analysis.xlsx"
    missing_report = args.output_dir / "files_missing_sf_error.csv"
    write_workbook(report, workbook)
    write_missing_error_report(report, missing_report)
    print(f"Failed files: {len(report.failed_files):,}")
    print(f"Missing sf__Error column: {len(report.missing_error_column):,}")
    print(f"Qualifying foreign-key errors: {report.qualifying_errors:,}")
    print(f"Referenced objects: {len(report.records_by_entity):,}")
    print(f"Workbook: {workbook}")
    print(f"Missing-column report: {missing_report}")
    if not args.no_copy:
        print(f"Copied failed files: {args.output_dir / 'failed_files'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
