"""Create a simple object-wise workbook for parked CampaignMembers."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def read_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        yield from csv.DictReader(handle)


def read_ids(path: Path, field: str) -> set[str]:
    return {
        str(row.get(field) or "").strip()
        for row in read_rows(path)
        if str(row.get(field) or "").strip()
    }


def style(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("folder", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    root = args.folder
    preload = root / "pre-load"
    source = preload / "FULL_CampaignMember_source_45k.csv"
    campaign_source = preload / "01_source_full" / "Campaign_R-54716_Campaign_3Aug.csv"
    inventory = preload / "03_existing_record"

    campaign_names = {
        str(row.get("Id") or "").strip(): str(row.get("Name") or "").strip()
        for row in read_rows(campaign_source)
        if str(row.get("Id") or "").strip()
    }
    campaign_ids = read_ids(inventory / "Campaign.csv", "Legacy_ID__c")
    contact_ids = read_ids(inventory / "Contact.csv", "Legacy_ID__c")
    successful_ids = set()
    for path in (root / "post-load").rglob("CampaignMember_success_*.csv"):
        successful_ids.update(read_ids(path, "Legacy_ID__c"))

    missing_lead = []
    missing_contact = []
    missing_campaign = []
    for row in read_rows(source):
        member_id = str(row.get("Id") or "").strip()
        if member_id in successful_ids:
            continue
        campaign_id = str(row.get("CampaignId") or "").strip()
        campaign_name = campaign_names.get(campaign_id, "")
        lead_id = str(row.get("LeadId") or "").strip()
        contact_id = str(row.get("ContactId") or "").strip()
        common = [member_id, campaign_id, campaign_name]
        if lead_id:
            missing_lead.append([*common, lead_id])
        elif contact_id and contact_id not in contact_ids:
            missing_contact.append([*common, contact_id])
        if campaign_id and campaign_id not in campaign_ids:
            missing_campaign.append([member_id, campaign_id, campaign_name])

    wb = Workbook()
    lead = wb.active
    lead.title = "Missing_Lead"
    lead.append(["CampaignMemberId", "CampaignId", "CampaignName", "LeadId"])
    for item in missing_lead:
        lead.append(item)
    style(lead, [26, 24, 55, 24])

    contact = wb.create_sheet("Missing_Contact")
    contact.append(["CampaignMemberId", "CampaignId", "CampaignName", "ContactId"])
    for item in missing_contact:
        contact.append(item)
    style(contact, [26, 24, 55, 24])

    if missing_campaign:
        campaign = wb.create_sheet("Missing_Campaign")
        campaign.append(["CampaignMemberId", "CampaignId", "CampaignName"])
        for item in missing_campaign:
            campaign.append(item)
        style(campaign, [26, 24, 55])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"missing_lead={len(missing_lead):,}")
    print(f"missing_contact={len(missing_contact):,}")
    print(f"missing_campaign={len(missing_campaign):,}")
    print(f"output={args.output}")


if __name__ == "__main__":
    main()
