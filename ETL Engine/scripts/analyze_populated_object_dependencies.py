from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SYSTEM_USER_FIELDS = {"CreatedById", "LastModifiedById", "OwnerId"}
STANDARD_REFERENCES = {
    "Task": {
        "WhoId": ("Contact", "Lead"),
        "WhatId": ("Account", "Asset", "Campaign", "Case", "Contract", "Opportunity", "Product2", "Solution"),
        "OwnerId": ("User", "Group"),
        "CreatedById": ("User",),
        "LastModifiedById": ("User",),
        "RecordTypeId": ("RecordType",),
    }
}


def csv_counts(path: Path) -> tuple[int, Counter[str]]:
    raw = path.read_bytes()
    encoding = "utf-8-sig"
    try:
        raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "cp1252"
    counts: Counter[str] = Counter()
    rows = 0
    with path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.DictReader(handle)
        row_iter = reader
        if len(reader.fieldnames or []) == 1 and "," in (reader.fieldnames or [""])[0]:
            # Some Salesforce exports wrap each complete CSV row in one quoted
            # outer cell. Parse that outer cell, then parse its inner CSV row.
            handle.seek(0)
            outer_rows = csv.reader(handle)
            inner_rows = (next(csv.reader([outer[0]])) for outer in outer_rows if outer)
            header = next(inner_rows, [])
            row_iter = (dict(zip(header, values)) for values in inner_rows)
        for row in row_iter:
            rows += 1
            for name, value in row.items():
                if name and value is not None and str(value).strip():
                    counts[name.strip()] += 1
    return rows, counts


def parse_schema(schema_root: Path, field_scope: set[str] | None = None) -> tuple[set[str], dict[str, dict[str, dict]]]:
    objects = {path.name for path in schema_root.iterdir() if path.is_dir()}
    fields: dict[str, dict[str, dict]] = defaultdict(dict)
    ns = {"m": "http://soap.sforce.com/2006/04/metadata"}
    for object_name in objects if field_scope is None else objects & field_scope:
        field_dir = schema_root / object_name / "fields"
        if not field_dir.is_dir():
            continue
        for path in field_dir.glob("*.field-meta.xml"):
            root = ET.parse(path).getroot()
            full_name = root.findtext("m:fullName", namespaces=ns) or path.name.removesuffix(".field-meta.xml")
            field_type = root.findtext("m:type", namespaces=ns) or ""
            references = tuple(
                item.text.strip() for item in root.findall("m:referenceTo", ns) if item.text and item.text.strip()
            )
            fields[object_name][full_name] = {
                "type": field_type,
                "references": references,
                "relationship_type": "Master-Detail" if field_type == "MasterDetail" else "Lookup",
                "required": field_type == "MasterDetail" or (
                    root.findtext("m:required", namespaces=ns) or "false"
                ).lower() == "true",
            }
    for object_name, refs in STANDARD_REFERENCES.items():
        for field_name, targets in refs.items():
            existing = fields[object_name].get(field_name)
            if existing and existing.get("references"):
                continue
            fields[object_name][field_name] = {
                "type": "Lookup", "references": targets, "relationship_type": "Lookup", "required": False,
            }
    return objects, fields


def strongly_connected_groups(nodes: list[str], edges: list[tuple[str, str]]) -> list[list[str]]:
    graph = {node: set() for node in nodes}
    for child, parent in edges:
        graph[child].add(parent)
    index = 0
    stack: list[str] = []
    on_stack: set[str] = set()
    indexes: dict[str, int] = {}
    low: dict[str, int] = {}
    result: list[list[str]] = []

    def visit(node: str) -> None:
        nonlocal index
        indexes[node] = low[node] = index
        index += 1
        stack.append(node)
        on_stack.add(node)
        for other in graph[node]:
            if other not in indexes:
                visit(other)
                low[node] = min(low[node], low[other])
            elif other in on_stack:
                low[node] = min(low[node], indexes[other])
        if low[node] == indexes[node]:
            group = []
            while True:
                member = stack.pop()
                on_stack.remove(member)
                group.append(member)
                if member == node:
                    break
            result.append(sorted(group))

    for node in sorted(nodes):
        if node not in indexes:
            visit(node)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--schema-root", type=Path, required=True)
    parser.add_argument("--objects-file", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    requested = [line.strip() for line in args.objects_file.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    requested_set = set(requested)
    schema_objects, schema_fields = parse_schema(args.schema_root, requested_set)

    files_by_object: dict[str, list[Path]] = defaultdict(list)
    for path in args.data_root.rglob("*.csv"):
        files_by_object[path.parent.name].append(path)

    population: dict[str, Counter[str]] = defaultdict(Counter)
    source_rows = []
    for object_name in requested:
        for path in files_by_object.get(object_name, []):
            rows, counts = csv_counts(path)
            population[object_name].update(counts)
            source_rows.append((object_name, str(path), rows, len(counts)))

    internal = []
    external = []
    excluded = []
    blank = []
    missing_columns = []
    for object_name in requested:
        for field_name, meta in sorted(schema_fields.get(object_name, {}).items()):
            if meta["type"] not in {"Lookup", "MasterDetail"}:
                continue
            count = population[object_name].get(field_name, 0)
            targets = meta["references"]
            row = (object_name, field_name, ", ".join(targets), meta["relationship_type"], count, "Yes" if meta["required"] else "No")
            if count == 0:
                blank.append(row)
            elif field_name in SYSTEM_USER_FIELDS or set(targets).issubset({"User", "Group"}):
                excluded.append(row + ("System/User relationship",))
            elif len(targets) != 1:
                external.append(row + ("Polymorphic; resolve by ID prefix/value",))
            elif targets[0] in requested_set:
                internal.append(row)
            else:
                reason = "Target exists in supplied schema; expected pre-existing in target UAT" if targets[0] in schema_objects else "DEPENDENCY TARGET MISSING FROM SUPPLIED SCHEMA"
                external.append(row + (reason,))

        for field_name, count in sorted(population[object_name].items()):
            if (field_name.endswith("Id") or field_name.endswith("__c")) and field_name not in schema_fields.get(object_name, {}) and field_name not in {"Id"}:
                missing_columns.append((object_name, field_name, count, "Not represented by a field metadata file; may be standard/non-relationship/formula"))

    edges = [(row[0], row[2]) for row in internal]
    components = strongly_connected_groups(requested, edges)
    component_by_object = {member: tuple(group) for group in components for member in group}
    parent_components = defaultdict(set)
    for child, parent in edges:
        if component_by_object[child] != component_by_object[parent]:
            parent_components[component_by_object[child]].add(component_by_object[parent])
    levels: dict[tuple[str, ...], int] = {}
    def level(group: tuple[str, ...]) -> int:
        if group not in levels:
            levels[group] = 1 + max((level(parent) for parent in parent_components[group]), default=0)
        return levels[group]
    for group in components:
        level(tuple(group))
    ordered = sorted((tuple(group) for group in components), key=lambda group: (levels[group], group))

    wb = Workbook()
    ws = wb.active
    ws.title = "ObjectHierarchy"
    ws.append(["MigrationOrder", "Wave", "ObjectName", "CircularGroup", "GroupMembers", "InternalParents", "MigrationInstruction"])
    order = 0
    parents = defaultdict(set)
    self_edges = set()
    for child, parent in edges:
        parents[child].add(parent)
        if child == parent:
            self_edges.add(child)
    for group in ordered:
        circular = len(group) > 1 or any(member in self_edges for member in group)
        for member in group:
            order += 1
            ws.append([order, levels[group], member, "Yes" if circular else "No", ", ".join(group), ", ".join(sorted(parents[member])), "Group in one pipeline; deferred lookup pass 2" if circular else "Load after internal parents"])

    sheets = {
        "InternalDependencies": (["Object", "Field", "ParentObject", "Type", "PopulatedRows", "Required"], internal),
        "ExternalTargetLookups": (["Object", "Field", "TargetObject", "Type", "PopulatedRows", "Required", "Classification"], external),
        "ExcludedSystemUser": (["Object", "Field", "TargetObject", "Type", "PopulatedRows", "Required", "Reason"], excluded),
        "BlankRelationships": (["Object", "Field", "TargetObject", "Type", "PopulatedRows", "Required"], blank),
        "SchemaValidation": (["Object", "InSchema", "HasDataFolder", "CSVFileCount"], [(name, "Yes" if name in schema_objects else "No", "Yes" if files_by_object.get(name) else "No", len(files_by_object.get(name, []))) for name in requested]),
        "UnclassifiedColumns": (["Object", "Column", "PopulatedRows", "Reason"], missing_columns),
        "SourceFiles": (["Object", "File", "Rows", "PopulatedColumnCount"], source_rows),
    }
    for title, (header, rows) in sheets.items():
        sheet = wb.create_sheet(title)
        sheet.append(header)
        for row in rows:
            sheet.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 70)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"Workbook: {args.output}")
    print(f"Requested objects: {len(requested)}; internal populated dependencies: {len(internal)}; external populated lookups: {len(external)}; excluded system/User: {len(excluded)}")
    print(f"Circular groups: {[list(group) for group in ordered if len(group) > 1 or any(member in self_edges for member in group)]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
