from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_RECORD_TYPE_ID = "012120000019jwqAAA"
DEVELOPER_NAME = "DS_IT"
TARGET_RECORD_TYPE_ID = "012S9000005vykZIAQ"


def autosize(sheet, maximum: int = 70) -> None:
    for column in range(1, sheet.max_column + 1):
        width = max(
            (len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), maximum)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)


def create_concise_report(existing: Path, output: Path, manifest_path: Path) -> None:
    """Replace the diagnostic workbook with a concise, linked report."""
    previous = load_workbook(existing, read_only=True)
    failure_rows = list(previous["Failure Summary"].iter_rows(min_row=2, values_only=True))
    previous.close()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Object", "Record received", "Prepared for reload", "Status", "Notes", "Primary data link"])
    summary.append([
        "GE_Installed_Base__c", manifest["input_rows"], manifest["output_rows"], "V3 PREPARED - REVALIDATE",
        "User and RecordType corrections applied. Account/date findings remain and require confirmation before claiming load-ready.",
        "Open v3 prepared CSV",
    ])
    summary["F2"].hyperlink = Path(manifest["output"]).as_uri()
    summary["F2"].style = "Hyperlink"

    changes = workbook.create_sheet("Applied Changes")
    changes.append(["Category", "Field", "Affected cells", "From", "To", "Evidence / notes"])
    for field in ("OwnerId", "CreatedById", "LastModifiedById"):
        originals = ", ".join(manifest["missing_user_values_by_field"].get(field, []))
        changes.append(["Fallback", field, manifest["changes_by_field"].get(field, 0), originals,
                        manifest["user_fallback_id"], "Only confirmed missing User values changed; valid values and blanks retained."])
    rt = manifest["record_type_mapping"]
    changes.append(["Transformation", "RecordTypeId", manifest["changes_by_field"].get("RecordTypeId", 0),
                    rt["source_id"], rt["developer_name"], f"Loader mapping: {rt['developer_name']} -> {rt['target_id']}"])
    changes.append(["Total", "All", manifest["total_cell_changes"], "", "", "Full row-level evidence is linked on Data Links."])

    outstanding = workbook.create_sheet("Outstanding")
    outstanding.append(["Field", "Affected rows in analyzed reject sample", "Finding", "Required next action"])
    for field, reason, count in failure_rows:
        if field in {"OwnerId", "CreatedById", "LastModifiedById", "RecordTypeId"}:
            continue
        action = "Provide/confirm Account fallback or mapping" if field == "Account__c" else "Confirm DD-MM-YYYY interpretation before date normalization"
        outstanding.append([field, count, reason, action])

    guidance = workbook.create_sheet("Loader Guidance")
    guidance.append(["Field", "Lookup-resolution selection", "Configuration"])
    guidance.append(["OwnerId", "External IDs by Object", "Add mapping: prefix 005; companion blank; User; Legacy_ID__c; Apply mapping"])
    guidance.append(["CreatedById", "External ID", "User.Legacy_ID__c"])
    guidance.append(["LastModifiedById", "External ID", "User.Legacy_ID__c"])
    guidance.append(["RecordTypeId", "RecordType mapping", f"{rt['developer_name']} -> {rt['target_id']}"])

    links = workbook.create_sheet("Data Links")
    links.append(["Artifact", "Purpose", "Status", "Path / link"])
    paths = [
        ("V2 source CSV", "Authoritative input; 28,555 rows", Path(manifest["source"])),
        ("V3 prepared CSV", "Reload input; 28,555 rows", Path(manifest["output"])),
        ("V3 transformation manifest", "Counts, mappings, and provenance", manifest_path.resolve()),
        ("V3 full fallback audit", "All 5,785 changed cells", Path(manifest["audit_file"])),
        ("Waygate_UAT DS_IT metadata", "Retrieved target RecordType evidence", Path(r"C:\VS Code\BH Dev\force-app\main\default\objects\GE_Installed_Base__c\recordTypes\DS_IT.recordType-meta.xml")),
        ("Original rejected CSV", "1,000-row failure evidence used for analysis", Path(manifest["rejected_evidence"])),
    ]
    for name, purpose, path in paths:
        exists = path.is_file()
        links.append([name, purpose, "Available" if exists else "Missing / moved", str(path)])
        if exists:
            cell = links.cell(links.max_row, 4)
            cell.hyperlink = path.resolve().as_uri()
            cell.style = "Hyperlink"

    notes = workbook.create_sheet("Reporting Notes")
    notes.append(["Topic", "Report note"])
    notes.append(["Scope", "This workbook reports the v2 failure analysis and the confirmed work used to create v3."])
    notes.append(["Evidence", "Fallback counts come from the v3 row-level audit, not merely from validation findings."])
    notes.append(["Rejected sample", "The analyzed rejected CSV contained 1,000 rows and is now missing from Downloads; its summarized findings are retained here."])
    notes.append(["Readiness", "V3 must be revalidated. Outstanding Account and date findings are intentionally not reported as resolved."])

    for sheet in workbook.worksheets:
        autosize(sheet)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main(source: Path, output: Path, manifest_path: Path | None = None, audit_path: Path | None = None) -> None:
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))

    issue_counts: Counter[tuple[str, str]] = Counter()
    row_details = []
    for row in rows:
        issues = [value.strip() for value in str(row.get("__validation_errors") or "").split(" | ") if value.strip()]
        for issue in issues:
            field, separator, reason = issue.partition(": ")
            issue_counts[(field if separator else "General", reason if separator else issue)] += 1
        row_details.append([
            row.get("__source_row"),
            row.get("Legacy_ID__c"),
            row.get("Name"),
            row.get("RecordTypeId"),
            DEVELOPER_NAME if row.get("RecordTypeId") == SOURCE_RECORD_TYPE_ID else "UNMAPPED",
            DEVELOPER_NAME if row.get("RecordTypeId") == SOURCE_RECORD_TYPE_ID else "",
            TARGET_RECORD_TYPE_ID if row.get("RecordTypeId") == SOURCE_RECORD_TYPE_ID else "",
            "REPLACE SOURCE ID WITH DEVELOPER NAME" if row.get("RecordTypeId") == SOURCE_RECORD_TYPE_ID else "REVIEW",
            len(issues),
            row.get("__validation_errors"),
        ])

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Generated UTC", datetime.now(timezone.utc).isoformat()])
    summary.append(["Source rejected CSV", str(source)])
    summary.append(["Object", "GE_Installed_Base__c"])
    summary.append(["Rejected rows", len(rows)])
    summary.append(["Rows with source RecordType DS_IT", sum(row[3] == SOURCE_RECORD_TYPE_ID for row in row_details)])
    summary.append(["Source RecordType ID", SOURCE_RECORD_TYPE_ID])
    summary.append(["RecordType DeveloperName", DEVELOPER_NAME])
    summary.append(["Waygate_UAT RecordType ID", TARGET_RECORD_TYPE_ID])
    summary.append(["CSV transformation", f"RecordTypeId: {SOURCE_RECORD_TYPE_ID} -> {DEVELOPER_NAME}"])
    summary.append(["Pipeline RecordType mapping", f"{DEVELOPER_NAME} -> {TARGET_RECORD_TYPE_ID}"])
    summary.append(["Important", "RecordType replacement resolves only RecordType errors; other listed validation failures remain."])

    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path else None
    audit = json.loads(audit_path.read_text(encoding="utf-8")) if audit_path else []
    if manifest:
        summary.append(["v3 output", manifest["output"]])
        summary.append(["v3 rows / columns", f"{manifest['output_rows']:,} rows / {manifest['output_columns']:,} columns"])
        summary.append(["v3 total transformed cells", manifest["total_cell_changes"]])
        summary.append(["User fallback", manifest["user_fallback_id"]])

    mapping = workbook.create_sheet("RecordType Transformation")
    mapping.append(["Object", "Source Org", "Source RecordType ID", "DeveloperName / CSV value", "Source Active", "Target Org", "Target RecordType ID / mapping result", "Target Active", "Action"])
    mapping.append(["GE_Installed_Base__c", "BH_dev", SOURCE_RECORD_TYPE_ID, DEVELOPER_NAME, "Yes", "Waygate_UAT", TARGET_RECORD_TYPE_ID, "Yes", "Replace source ID with DS_IT; configure DS_IT -> target ID"])

    failures = workbook.create_sheet("Failure Summary")
    failures.append(["Field", "Failure reason", "Affected rows"])
    for (field, reason), count in sorted(issue_counts.items(), key=lambda item: (-item[1], item[0])):
        failures.append([field, reason, count])

    details = workbook.create_sheet("Rejected Row Remediation")
    details.append(["Source row", "Legacy_ID__c", "Name", "Original RecordTypeId", "Source DeveloperName", "Transformed CSV value", "Resolved target RecordTypeId", "RecordType action", "Validation issue count", "All validation errors"])
    for detail in row_details:
        details.append(detail)
    for cell in details["J"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    if manifest:
        execution = workbook.create_sheet("v3 Execution Summary")
        execution.append(["Item", "Value"])
        for label, key in (
            ("Created UTC", "created_at"), ("Input v2", "source"), ("Rejected evidence", "rejected_evidence"),
            ("Output v3", "output"), ("Input rows", "input_rows"), ("Output rows", "output_rows"),
            ("Input columns", "input_columns"), ("Output columns", "output_columns"),
            ("Total changed cells", "total_cell_changes"), ("Fallback audit", "audit_file"),
        ):
            execution.append([label, manifest.get(key)])
        execution.append(["Result", "Row and column counts retained; only confirmed User fallback values and RecordType DS_IT values were transformed."])

        fallback = workbook.create_sheet("User Fallback Mapping")
        fallback.append(["Field", "Referenced object", "Missing source value", "Replacement", "Changed cells", "Resolution policy"])
        for field, missing_values in manifest["missing_user_values_by_field"].items():
            for value in missing_values:
                fallback.append([field, "User", value, manifest["user_fallback_id"], manifest["changes_by_field"].get(field, 0), "User.Legacy_ID__c"])

        full_audit = workbook.create_sheet("v3 Full Change Audit")
        audit_headers = ["source_row", "record_id", "field", "original", "replacement", "resolved_target_id", "reason"]
        full_audit.append(audit_headers)
        for item in audit:
            full_audit.append([item.get(name, "") for name in audit_headers])

        outstanding = workbook.create_sheet("Outstanding Before Reload")
        outstanding.append(["Field", "Failure reason", "Affected rejected rows", "Status / required action"])
        for (field, reason), count in sorted(issue_counts.items(), key=lambda item: (-item[1], item[0])):
            if field in {"OwnerId", "CreatedById", "LastModifiedById", "RecordTypeId"}:
                continue
            action = "Supply an Account fallback or valid mapping" if field == "Account__c" else "Normalize date after confirming source date convention"
            outstanding.append([field, reason, count, action])

        guidance = workbook.create_sheet("Loader Guidance")
        guidance.append(["Field", "Select in loader", "Configuration"])
        guidance.append(["OwnerId", "External IDs by Object", "Add object mapping: Salesforce ID prefix 005; companion blank; object User; target external ID Legacy_ID__c; Apply mapping"])
        guidance.append(["CreatedById", "External ID", "User.Legacy_ID__c"])
        guidance.append(["LastModifiedById", "External ID", "User.Legacy_ID__c"])
        guidance.append(["RecordTypeId", "RecordType mapping", f"{DEVELOPER_NAME} -> {TARGET_RECORD_TYPE_ID}"])

    for sheet in workbook.worksheets:
        autosize(sheet)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Created {output}")
    print(f"Rows: {len(rows)}; distinct failure types: {len(issue_counts)}")


if __name__ == "__main__":
    if len(sys.argv) not in {3, 5}:
        raise SystemExit("Usage: create_ge_installed_base_failure_workbook.py REJECTED.csv OUTPUT.xlsx [V3_MANIFEST.json V3_AUDIT.json]")
    main(
        Path(sys.argv[1]), Path(sys.argv[2]),
        Path(sys.argv[3]) if len(sys.argv) == 5 else None,
        Path(sys.argv[4]) if len(sys.argv) == 5 else None,
    )
