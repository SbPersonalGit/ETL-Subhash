"""Cross-verify populated hierarchy dependencies and wave ordering."""

from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
load_dotenv(PROJECT_ROOT / ".env", override=True)

from core.connector import connector_from_env

WORKBOOK = PROJECT_ROOT / "output" / "uat_to_prod_populated_hierarchy.xlsx"
COUNTS = PROJECT_ROOT / "output" / "legacy_id_soql_counts_updated.csv"


def main() -> int:
    wb = load_workbook(WORKBOOK, data_only=True)

    # Wave / record counts
    waves_ws = wb["Waves"]
    wave_headers = [c.value for c in waves_ws[1]]
    wave_by_object: dict[str, int] = {}
    count_by_object: dict[str, int] = {}
    for row in waves_ws.iter_rows(min_row=2, values_only=True):
        obj = row[wave_headers.index("ObjectName")]
        wave_label = str(row[wave_headers.index("Wave")] or "")
        level = int(row[wave_headers.index("DependencyLevel")] or 0)
        record_count = row[wave_headers.index("RecordCount")]
        wave_by_object[obj] = level
        count_by_object[obj] = int(record_count or 0)

    # Expected counts from source CSV
    expected = {
        r["sObject"]: int(r["RecordCount"])
        for r in csv.DictReader(COUNTS.open(encoding="utf-8"))
    }

    # Internal populated deps
    deps_ws = wb["PopulatedInternalDeps"]
    dep_headers = [c.value for c in deps_ws[1]]
    deps = []
    for row in deps_ws.iter_rows(min_row=2, values_only=True):
        deps.append(
            {
                "object": row[dep_headers.index("ObjectName")],
                "parent": row[dep_headers.index("ParentObject")],
                "field": row[dep_headers.index("ParentIdField")],
                "type": row[dep_headers.index("RelationshipType")],
                "populated": int(row[dep_headers.index("PopulatedRows")] or 0),
            }
        )

    print("=== 1) RecordCount cross-check vs Legacy_ID SOQL CSV ===")
    count_mismatches = []
    for obj, count in count_by_object.items():
        if expected.get(obj) != count:
            count_mismatches.append((obj, count, expected.get(obj)))
    print(f"objects={len(count_by_object)} mismatches={len(count_mismatches)}")
    for item in count_mismatches[:10]:
        print(" MISMATCH", item)

    print("\n=== 2) Wave order vs populated internal dependencies ===")
    # Child wave must be >= parent wave. Equal only allowed for cycles/self.
    order_violations = []
    same_wave_edges = []
    for dep in deps:
        child, parent = dep["object"], dep["parent"]
        if child not in wave_by_object or parent not in wave_by_object:
            order_violations.append({**dep, "reason": "missing_from_waves"})
            continue
        child_wave = wave_by_object[child]
        parent_wave = wave_by_object[parent]
        if child_wave < parent_wave:
            order_violations.append(
                {
                    **dep,
                    "reason": "child_before_parent",
                    "child_wave": child_wave,
                    "parent_wave": parent_wave,
                }
            )
        elif child_wave == parent_wave:
            same_wave_edges.append({**dep, "wave": child_wave})

    print(f"populated_internal_deps={len(deps)}")
    print(f"order_violations={len(order_violations)}")
    for item in order_violations[:20]:
        print(" VIOLATION", item)
    print(f"same_wave_edges(cycle_or_self)={len(same_wave_edges)}")

    # Group same-wave by objects
    same_wave_pairs = sorted({(e["object"], e["parent"], e["field"]) for e in same_wave_edges})
    print("sample same-wave edges:")
    for item in same_wave_pairs[:15]:
        print(" ", item)

    print("\n=== 3) Live SOQL re-check of sample edges (populated still > 0) ===")
    connector = connector_from_env("TARGET")
    connector.connect()

    # Prefer high-volume / important chain samples
    samples = []
    preferred = [
        ("Contact", "AccountId", "Account"),
        ("Opportunity", "AccountId", "Account"),
        ("OpportunityLineItem", "OpportunityId", "Opportunity"),
        ("Case", "AccountId", "Account"),
        ("Case", "ContactId", "Contact"),
        ("EmailMessage", "ParentId", "Case"),
        ("Contract", "AccountId", "Account"),
        ("Deal_Machine__kav", None, "GE_OG_Product_Hierarchy1__c"),
        ("SVMXC__Service_Order_Line__c", None, "SVMXC__Service_Order__c"),
        ("SFPRM_DS_Invoice__c", None, "SFPRM_DS_Order__c"),
    ]
    by_key = {(d["object"], d["parent"]): d for d in deps}
    by_obj_parent_field = {(d["object"], d["field"], d["parent"]): d for d in deps}

    for obj, field, parent in preferred:
        if field:
            dep = by_obj_parent_field.get((obj, field, parent)) or by_key.get((obj, parent))
        else:
            dep = by_key.get((obj, parent))
            if dep:
                field = dep["field"]
        if not dep:
            samples.append({"object": obj, "parent": parent, "field": field, "status": "NOT_IN_DEPS"})
            continue
        field = dep["field"]
        soql = f"SELECT COUNT(Id) cnt FROM {obj} WHERE {field} != null"
        try:
            rows = connector.query(soql)
            live = int(rows[0].get("cnt") or 0)
            status = "OK" if live > 0 else "NOW_ZERO"
            if live != dep["populated"]:
                status = "COUNT_CHANGED"
            samples.append(
                {
                    "object": obj,
                    "field": field,
                    "parent": parent,
                    "workbook_populated": dep["populated"],
                    "live_populated": live,
                    "status": status,
                    "child_wave": wave_by_object.get(obj),
                    "parent_wave": wave_by_object.get(parent),
                }
            )
        except Exception as exc:
            samples.append(
                {
                    "object": obj,
                    "field": field,
                    "parent": parent,
                    "status": f"ERROR: {exc}",
                }
            )

    for sample in samples:
        print(sample)

    print("\n=== 4) Orphan-ish check: child lookup values exist on parent (sample) ===")
    orphan_checks = [
        ("Contact", "AccountId", "Account"),
        ("OpportunityLineItem", "OpportunityId", "Opportunity"),
        ("EmailMessage", "ParentId", "Case"),
        ("SFPRM_DS_Invoice__c", None, "SFPRM_DS_Order__c"),
    ]
    for obj, field, parent in orphan_checks:
        dep = by_obj_parent_field.get((obj, field, parent)) if field else by_key.get((obj, parent))
        if not dep:
            # find any field to that parent
            dep = next((d for d in deps if d["object"] == obj and d["parent"] == parent), None)
        if not dep:
            print({"object": obj, "parent": parent, "status": "NOT_IN_DEPS"})
            continue
        field = dep["field"]
        # Count child rows whose lookup is populated but parent Id not found
        soql = (
            f"SELECT COUNT(Id) cnt FROM {obj} "
            f"WHERE {field} != null AND {field} NOT IN (SELECT Id FROM {parent})"
        )
        try:
            rows = connector.query(soql)
            orphans = int(rows[0].get("cnt") or 0)
            print(
                {
                    "edge": f"{obj}.{field} -> {parent}",
                    "orphan_count": orphans,
                    "status": "OK" if orphans == 0 else "ORPHANS_FOUND",
                }
            )
        except Exception as exc:
            # Semi-join not always allowed; fallback sample compare
            print({"edge": f"{obj}.{field} -> {parent}", "status": f"SEMIJOIN_FAILED: {exc}"})
            try:
                child_ids = connector.query(f"SELECT {field} FROM {obj} WHERE {field} != null LIMIT 200")
                values = [r[field] for r in child_ids if r.get(field)]
                if not values:
                    print({"edge": f"{obj}.{field} -> {parent}", "status": "NO_SAMPLE_VALUES"})
                    continue
                id_list = ",".join(f"'{v}'" for v in values[:100])
                parents = connector.query(f"SELECT Id FROM {parent} WHERE Id IN ({id_list})")
                found = {r["Id"] for r in parents}
                missing = [v for v in values[:100] if v not in found]
                print(
                    {
                        "edge": f"{obj}.{field} -> {parent}",
                        "sampled": min(100, len(values)),
                        "missing_in_parent": len(missing),
                        "status": "OK" if not missing else "ORPHANS_IN_SAMPLE",
                    }
                )
            except Exception as exc2:
                print({"edge": f"{obj}.{field} -> {parent}", "status": f"SAMPLE_FAILED: {exc2}"})

    summary = {
        "recordcount_mismatches": len(count_mismatches),
        "order_violations": len(order_violations),
        "same_wave_edges": len(same_wave_edges),
        "sample_edge_statuses": [s.get("status") for s in samples],
    }
    out = PROJECT_ROOT / "output" / "hierarchy_dependency_crosscheck.summary.json"
    out.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SUMMARY ===")
    print(json.dumps(summary, indent=2))
    return 0 if not order_violations and not count_mismatches else 1


if __name__ == "__main__":
    raise SystemExit(main())
