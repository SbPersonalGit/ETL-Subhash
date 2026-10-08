#!/usr/bin/env python3
"""Apply lookup fallback IDs from a rejected prevalidation CSV.

Creates a new versioned file (v2, v3, ...) under Input Files / {Object API}
and a fallback audit CSV next to the rejected file.
Never renames or overwrites the source vN file.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.flat_file import open_csv  # noqa: E402
from core.workspace_layout import INPUT_FILES_DIR  # noqa: E402

DEFAULT_DROP_ROOT = INPUT_FILES_DIR

REFERENCE_ERROR_RE = re.compile(
    r"reference '([^']+)' not found in ([A-Za-z0-9_]+)\.External_Id__c",
    re.IGNORECASE,
)
VERSIONED_NAME_RE = re.compile(
    r"^(?P<object>.+)\.v(?P<version>\d+)(?P<suffix>_rejected|_clean)?$",
    re.IGNORECASE,
)
USER_COLUMNS = {
    "ownerid": "User",
    "createdbyid": "User",
    "lastmodifiedbyid": "User",
}
AUDIT_HEADERS = [
    "SourceFile",
    "UpdatedFile",
    "SourceRow",
    "Id",
    "Name",
    "PromotionCode",
    "PromotionName",
    "Column",
    "ReferencedObject",
    "OriginalValue",
    "UpdatedValue",
    "Reason",
]


def _pick(row: dict, *names: str) -> str:
    lower = {str(key).casefold(): key for key in row}
    for name in names:
        key = lower.get(name.casefold())
        if key is not None:
            return str(row.get(key) or "").strip()
    return ""


def _column(row: dict, name: str) -> str | None:
    lower = {str(key).casefold(): key for key in row}
    return lower.get(name.casefold())


def parse_versioned_stem(path: Path) -> tuple[str, int, str]:
    match = VERSIONED_NAME_RE.match(path.stem)
    if not match:
        raise ValueError(
            f"{path.name}: expected a versioned name like Object__c.v1.csv or "
            "Object__c.v1_rejected.csv"
        )
    return match.group("object"), int(match.group("version")), match.group("suffix") or ""


def load_fallback_ids(xlsx_path: Path) -> dict[str, str]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"{xlsx_path}: fallback workbook is empty")
    headers = [str(cell or "").strip() for cell in rows[0]]
    header_map = {name.casefold(): index for index, name in enumerate(headers) if name}
    object_idx = header_map.get("object")
    id_idx = header_map.get("external_id__c", header_map.get("fallbackid"))
    if object_idx is None or id_idx is None:
        raise ValueError(
            f"{xlsx_path}: expected columns Object and External_Id__c (or FallbackId)"
        )
    fallbacks: dict[str, str] = {}
    for row in rows[1:]:
        if not row:
            continue
        object_name = str(row[object_idx] or "").strip()
        fallback_id = str(row[id_idx] or "").strip()
        if object_name and fallback_id:
            fallbacks[object_name] = fallback_id
    if not fallbacks:
        raise ValueError(f"{xlsx_path}: no fallback IDs found")
    return fallbacks


def missing_ids_from_rejected(rejected_path: Path) -> dict[str, set[str]]:
    missing: dict[str, set[str]] = defaultdict(set)
    with open_csv(rejected_path) as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            blobs = [str(row.get("__validation_errors") or "")]
            for key, value in row.items():
                if str(key).startswith("__error_") and str(value or "").strip():
                    blobs.append(str(value))
            for blob in blobs:
                for match in REFERENCE_ERROR_RE.finditer(blob):
                    missing[match.group(2)].add(match.group(1))
            for folded, object_name in USER_COLUMNS.items():
                error_key = _column(row, f"__error_{folded}")
                if not error_key or not str(row.get(error_key) or "").strip():
                    continue
                data_key = _column(row, folded)
                if data_key:
                    value = str(row.get(data_key) or "").strip()
                    if value:
                        missing[object_name].add(value)
    return {object_name: ids for object_name, ids in missing.items() if ids}


def resolve_source_csv(rejected_path: Path, object_name: str, version: int, source: Path | None) -> Path:
    if source:
        path = source
        if not path.exists():
            raise FileNotFoundError(f"source CSV not found: {path}")
        return path
    filename = f"{object_name}.v{version}.csv"
    candidates = [
        DEFAULT_DROP_ROOT / object_name / filename,
        rejected_path.parent / filename,
    ]
    for path in candidates:
        if path.exists():
            return path
    searched = " ; ".join(str(path) for path in candidates)
    raise FileNotFoundError(f"could not find {filename}. Looked in: {searched}")


def next_version_path(source_path: Path, object_name: str, version: int) -> Path:
    return source_path.with_name(f"{object_name}.v{version + 1}.csv")


def audit_path(rejected_path: Path, object_name: str, version: int) -> Path:
    return rejected_path.with_name(f"{object_name}.V{version}.fallback.Updated.csv")


def apply_fallbacks(
    *,
    rejected_path: Path,
    fallbacks_path: Path,
    source_path: Path | None = None,
) -> dict:
    rejected_path = rejected_path.resolve()
    fallbacks_path = fallbacks_path.resolve()
    object_name, version, _suffix = parse_versioned_stem(rejected_path)
    source_path = resolve_source_csv(rejected_path, object_name, version, source_path)
    source_object, source_version, _ = parse_versioned_stem(source_path)
    if source_object.casefold() != object_name.casefold():
        raise ValueError(
            f"rejected object {object_name} does not match source {source_path.name}"
        )
    dest_path = next_version_path(source_path, object_name, source_version)
    if dest_path.resolve() == source_path.resolve():
        raise ValueError("refusing to overwrite the source CSV")
    report_path = audit_path(rejected_path, object_name, source_version)

    fallbacks = load_fallback_ids(fallbacks_path)
    missing = missing_ids_from_rejected(rejected_path)
    if not missing:
        raise ValueError(f"{rejected_path.name}: no lookup fallback errors were found")

    with open_csv(source_path) as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)
    if not fieldnames:
        raise ValueError(f"{source_path.name}: CSV has no header")

    audit: list[dict[str, str]] = []
    cells = 0
    row_hits = 0
    replaced_ids: Counter[str] = Counter()
    for index, row in enumerate(rows, start=2):
        changed = False
        rec_id = _pick(row, "Id", "External_Id__c")
        rec_name = _pick(row, "Name")
        promo_code = _pick(row, "APTS_PROMOTION_CODE__C", "APTS_Promotion_Code__c")
        promo_name = _pick(
            row, "APTS_PROMOTION_NAME__C", "APTS_Promotion_Name__c", "Name"
        )
        for key in list(row.keys()):
            object_name_for_col = USER_COLUMNS.get(str(key).casefold())
            if not object_name_for_col:
                continue
            fallback_id = fallbacks.get(object_name_for_col)
            if not fallback_id:
                continue
            original = str(row.get(key) or "").strip()
            if original and original in missing.get(object_name_for_col, set()) and original != fallback_id:
                row[key] = fallback_id
                cells += 1
                changed = True
                replaced_ids[original] += 1
                audit.append({
                    "SourceFile": source_path.name,
                    "UpdatedFile": dest_path.name,
                    "SourceRow": str(index),
                    "Id": rec_id,
                    "Name": rec_name,
                    "PromotionCode": promo_code,
                    "PromotionName": promo_name,
                    "Column": key,
                    "ReferencedObject": object_name_for_col,
                    "OriginalValue": original,
                    "UpdatedValue": fallback_id,
                    "Reason": (
                        f"reference '{original}' not found in "
                        f"{object_name_for_col}.External_Id__c"
                    ),
                })
        if changed:
            row_hits += 1

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    with dest_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=AUDIT_HEADERS)
        writer.writeheader()
        writer.writerows(audit)

    unused = sorted(set(fallbacks) - set(missing))
    return {
        "object": object_name,
        "source": str(source_path),
        "updated": str(dest_path),
        "audit": str(report_path),
        "rows": len(rows),
        "rows_updated": row_hits,
        "cells_replaced": cells,
        "distinct_original_ids": len(replaced_ids),
        "missing_by_object": {key: len(value) for key, value in missing.items()},
        "fallbacks": fallbacks,
        "unused_fallback_objects": unused,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Create Object.vN+1.csv and Object.VN.fallback.Updated.csv from a "
            "rejected prevalidation file and KLG fallback IDs. Does not change v1."
        )
    )
    parser.add_argument("--rejected", required=True, type=Path, help="Rejected CSV path")
    parser.add_argument(
        "--fallbacks",
        required=True,
        type=Path,
        help="Excel with Object, External_Id__c (KLG Fallback Ids.xlsx)",
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=None,
        help="Optional source CSV (default: Input Files/{Object}/{Object}.vN.csv)",
    )
    args = parser.parse_args(argv)
    result = apply_fallbacks(
        rejected_path=args.rejected,
        fallbacks_path=args.fallbacks,
        source_path=args.source,
    )
    print(f"source     {result['source']}")
    print(f"updated    {result['updated']}")
    print(f"audit      {result['audit']}")
    print(
        f"rows       {result['rows_updated']} / {result['rows']}  "
        f"cells {result['cells_replaced']}  "
        f"ids {result['distinct_original_ids']}"
    )
    print(f"fallbacks  {result['fallbacks']}")
    if result["unused_fallback_objects"]:
        print(f"unused     {result['unused_fallback_objects']}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
