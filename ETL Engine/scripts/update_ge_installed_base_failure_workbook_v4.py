"""Rebuild the GE_Installed_Base failure/RecordType workbook for v4."""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def autosize(sheet, maximum: int = 70) -> None:
    for column in range(1, sheet.max_column + 1):
        width = max(
            (len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 200) + 1)),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), maximum)
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1 and sheet.max_column > 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)


def iter_issues(row: dict[str, str]) -> list[tuple[str, str]]:
    issues: list[tuple[str, str]] = []
    combined = str(row.get("__validation_errors") or "").strip()
    if combined:
        for part in combined.split(" | "):
            field, sep, reason = part.partition(": ")
            if sep:
                issues.append((field.strip(), reason.strip()))
            elif part.strip():
                issues.append(("General", part.strip()))
        return issues
    for key, value in row.items():
        if key.startswith("__error_") and str(value or "").strip():
            issues.append((key[len("__error_") :], str(value).strip()))
    return issues


def main(rejected: Path, output: Path, manifest_path: Path, audit_path: Path) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    audit = json.loads(audit_path.read_text(encoding="utf-8"))

    with rejected.open("r", encoding="utf-8-sig", newline="") as handle:
        rejected_rows = list(csv.DictReader(handle))

    issue_counts: Counter[tuple[str, str]] = Counter()
    rt_before: Counter[str] = Counter()
    remediation_rows = []
    for row in rejected_rows:
        issues = iter_issues(row)
        for field, reason in issues:
            issue_counts[(field, reason)] += 1
        original_rt = str(row.get("RecordTypeId") or "").strip()
        rt_before[original_rt] += 1
        mapped = {
            "DS_IT": ("DS_IT", "012S9000005vykZIAQ", "KEEP / MAP DEVELOPER NAME"),
            "012A0000000jpV9IAI": ("Inspection", "012S9000005vykrIAA", "REPLACE SOURCE ID WITH DEVELOPER NAME"),
            "012120000019jwqAAA": ("DS_IT", "012S9000005vykZIAQ", "REPLACE SOURCE ID WITH DEVELOPER NAME"),
            "Inspection": ("Inspection", "012S9000005vykrIAA", "KEEP / MAP DEVELOPER NAME"),
        }.get(original_rt, ("UNMAPPED", "", "REVIEW"))
        remediation_rows.append([
            row.get("__source_row") or "",
            row.get("Legacy_ID__c") or row.get("Id") or "",
            row.get("Name") or "",
            original_rt,
            mapped[0] if mapped[0] != "UNMAPPED" else "UNMAPPED",
            mapped[0] if mapped[0] != "UNMAPPED" else "",
            mapped[1],
            mapped[2],
            len(issues),
            " | ".join(f"{field}: {reason}" for field, reason in issues),
        ])

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Generated UTC", datetime.now(timezone.utc).isoformat()])
    summary.append(["Source rejected CSV", str(rejected)])
    summary.append(["Object", "GE_Installed_Base__c"])
    summary.append(["Rejected rows", len(rejected_rows)])
    summary.append(["Rows with RecordType DS_IT", rt_before.get("DS_IT", 0)])
    summary.append(["Rows with source Inspection ID 012A0000000jpV9IAI", rt_before.get("012A0000000jpV9IAI", 0)])
    summary.append(["Waygate_UAT DS_IT RecordType ID", "012S9000005vykZIAQ"])
    summary.append(["Waygate_UAT Inspection RecordType ID", "012S9000005vykrIAA"])
    summary.append(["CSV transformation", "012A0000000jpV9IAI -> Inspection; DS_IT retained as DeveloperName"])
    summary.append(["Pipeline RecordType mapping", "DS_IT -> 012S9000005vykZIAQ; Inspection -> 012S9000005vykrIAA"])
    summary.append(["User fallback", manifest.get("user_fallback_id")])
    summary.append(["Account fallback", manifest.get("account_fallback_id")])
    summary.append(["v4 output", manifest.get("output")])
    summary.append([
        "v4 rows / columns",
        f"{manifest.get('output_rows', 0):,} rows / {manifest.get('output_columns', 0):,} columns",
    ])
    summary.append(["v4 total transformed cells", manifest.get("total_cell_changes")])
    summary.append(["v4 fallback audit JSON", manifest.get("audit_file")])

    mapping = workbook.create_sheet("RecordType Transformation")
    mapping.append([
        "Object", "Source Org", "Source RecordType ID / value", "DeveloperName / CSV value",
        "Source Active", "Target Org", "Target RecordType ID / mapping result", "Target Active", "Action",
    ])
    mapping.append([
        "GE_Installed_Base__c", "BH_dev", "012A0000000jpV9IAI", "Inspection", "Yes",
        "Waygate_UAT", "012S9000005vykrIAA", "Yes",
        "Replace source Inspection ID with DeveloperName Inspection; configure Inspection -> target ID",
    ])
    mapping.append([
        "GE_Installed_Base__c", "BH_dev / prior v3", "DS_IT", "DS_IT", "Yes",
        "Waygate_UAT", "012S9000005vykZIAQ", "Yes",
        "Retain DeveloperName DS_IT; configure DS_IT -> target ID",
    ])

    failures = workbook.create_sheet("Failure Summary")
    failures.append(["Field", "Failure reason", "Affected rows"])
    for (field, reason), count in sorted(issue_counts.items(), key=lambda item: (-item[1], item[0])):
        failures.append([field, reason, count])

    details = workbook.create_sheet("Rejected Row Remediation")
    details.append([
        "Source row", "Legacy_ID__c", "Name", "Original RecordTypeId", "Source DeveloperName",
        "Transformed CSV value", "Resolved target RecordTypeId", "RecordType action",
        "Validation issue count", "All validation errors",
    ])
    # Keep workbook usable: first 2,000 rejected rows plus note.
    max_detail = 2000
    for detail in remediation_rows[:max_detail]:
        details.append(detail)
    if len(remediation_rows) > max_detail:
        details.append([
            "", "", f"... truncated {len(remediation_rows) - max_detail:,} additional rejected rows",
            "", "", "", "", "", "", f"See rejected CSV and v4 audit JSON for full detail",
        ])
    for cell in details["J"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    execution = workbook.create_sheet("v4 Execution Summary")
    execution.append(["Item", "Value"])
    for label, key in (
        ("Created UTC", "created_at"),
        ("Input v3", "source"),
        ("Rejected evidence", "rejected_evidence"),
        ("Output v4", "output"),
        ("Input rows", "input_rows"),
        ("Output rows", "output_rows"),
        ("Input columns", "input_columns"),
        ("Output columns", "output_columns"),
        ("Total changed cells", "total_cell_changes"),
        ("Fallback audit", "audit_file"),
        ("User fallback", "user_fallback_id"),
        ("Account fallback", "account_fallback_id"),
    ):
        execution.append([label, manifest.get(key)])
    execution.append([
        "Result",
        "Row/column counts retained; applied User/Account fallbacks, Inspection+DS_IT RecordType DeveloperName transforms, and DD-MM-YYYY -> YYYY-MM-DD normalization.",
    ])

    user_map = workbook.create_sheet("User Fallback Mapping")
    user_map.append(["Field", "Referenced object", "Missing source value", "Replacement", "Changed cells", "Resolution policy"])
    for field, missing_values in (manifest.get("missing_user_values_by_field") or {}).items():
        for value in missing_values:
            user_map.append([
                field, "User", value, manifest.get("user_fallback_id"),
                (manifest.get("changes_by_field") or {}).get(field, 0), "User.Legacy_ID__c",
            ])

    account_map = workbook.create_sheet("Account Fallback Mapping")
    account_map.append(["Field", "Referenced object", "Missing source value", "Replacement", "Changed cells", "Resolution policy"])
    for field, missing_values in (manifest.get("missing_account_values_by_field") or {}).items():
        for value in missing_values:
            account_map.append([
                field, "Account", value, manifest.get("account_fallback_id"),
                (manifest.get("changes_by_field") or {}).get(field, 0), "Account.Legacy_ID__c",
            ])

    change_summary = workbook.create_sheet("v4 Change Summary")
    change_summary.append(["Field", "Changed cells"])
    for field, count in sorted((manifest.get("changes_by_field") or {}).items(), key=lambda item: (-item[1], item[0])):
        change_summary.append([field, count])
    change_summary.append(["TOTAL", manifest.get("total_cell_changes")])

    # Sample of audit rows for Excel; full audit remains in JSON.
    full_audit = workbook.create_sheet("v4 Full Change Audit Sample")
    audit_headers = ["source_row", "record_id", "field", "original", "replacement", "resolved_target_id", "reason"]
    full_audit.append(audit_headers)
    sample_n = 5000
    for item in audit[:sample_n]:
        full_audit.append([item.get(name, "") for name in audit_headers])
    if len(audit) > sample_n:
        full_audit.append(["", "", "", "", "", "", f"... truncated; full audit has {len(audit):,} rows in JSON"])

    outstanding = workbook.create_sheet("Outstanding Before Reload")
    outstanding.append(["Field", "Failure reason", "Affected rejected rows", "Status / required action"])
    handled_fields = {
        "OwnerId", "CreatedById", "LastModifiedById", "GE_ES_Account_Manager__c",
        "RecordTypeId", "Account__c",
        *(manifest.get("date_fields_normalized") or []),
    }
    remaining = [
        ((field, reason), count)
        for (field, reason), count in issue_counts.items()
        if field not in handled_fields
    ]
    if not remaining:
        outstanding.append(["(none)", "All observed rejection classes were remediated in v4", 0, "Ready for reload / prevalidation"])
    else:
        for (field, reason), count in sorted(remaining, key=lambda item: (-item[1], item[0])):
            outstanding.append([field, reason, count, "Manual review required"])

    guidance = workbook.create_sheet("Loader Guidance")
    guidance.append(["Field", "Select in loader", "Configuration"])
    guidance.append([
        "OwnerId", "External IDs by Object",
        "Salesforce ID prefix 005; object User; target external ID Legacy_ID__c",
    ])
    guidance.append(["CreatedById", "External ID", "User.Legacy_ID__c"])
    guidance.append(["LastModifiedById", "External ID", "User.Legacy_ID__c"])
    guidance.append(["GE_ES_Account_Manager__c", "External ID", "User.Legacy_ID__c"])
    guidance.append(["Account__c", "External ID", "Account.Legacy_ID__c"])
    guidance.append(["RecordTypeId", "RecordType mapping", "DS_IT -> 012S9000005vykZIAQ"])
    guidance.append(["RecordTypeId", "RecordType mapping", "Inspection -> 012S9000005vykrIAA"])

    for sheet in workbook.worksheets:
        autosize(sheet)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Updated {output}")
    print(f"Rejected rows: {len(rejected_rows)}; audit rows: {len(audit)}; failure types: {len(issue_counts)}")


if __name__ == "__main__":
    if len(sys.argv) != 5:
        raise SystemExit(
            "Usage: update_ge_installed_base_failure_workbook_v4.py REJECTED.csv OUTPUT.xlsx MANIFEST.json AUDIT.json"
        )
    main(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), Path(sys.argv[4]))
