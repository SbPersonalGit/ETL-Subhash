"""Audit hierarchy objects and Legacy_ID__c availability in the target org."""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.connector import connector_from_env


def hierarchy_objects(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["ObjectHierarchy"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    column = headers.index("ObjectName")
    result = []
    seen = set()
    for row in rows:
        value = str(row[column] or "").strip() if len(row) > column else ""
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    workbook.close()
    return result


def style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in sheet.iter_rows(min_row=2):
        status = row[2].value if len(row) > 2 else ""
        fill = "C6EFCE" if status == "READY" else "FFEB9C" if status in {"FIELD_NOT_EXTERNAL_ID", "FIELD_NOT_WRITABLE"} else "FFC7CE"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=fill)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("hierarchy", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--target-profile", default=None)
    args = parser.parse_args()

    objects = hierarchy_objects(args.hierarchy)
    load_dotenv(PROJECT_ROOT / ".env", override=True)
    connector = connector_from_env("TARGET", profile=args.target_profile)
    connector.connect()
    global_objects = {entry["name"]: entry for entry in connector.describe_global()}

    rows = []
    for order, object_name in enumerate(objects, 1):
        if object_name not in global_objects:
            rows.append({
                "Order": order,
                "Object": object_name,
                "Status": "OBJECT_MISSING",
                "ObjectExists": "No",
                "LegacyFieldExists": "No",
                "FieldType": "",
                "Length": "",
                "ExternalId": "",
                "Unique": "",
                "Createable": "",
                "Updateable": "",
                "ObjectLabel": "",
                "Detail": "Object is not available in the connected target org.",
            })
            continue
        try:
            describe = connector.describe_object(object_name)
            field = next((item for item in describe.get("fields", []) if item.get("name", "").lower() == "legacy_id__c"), None)
            if not field:
                status = "LEGACY_FIELD_MISSING"
                detail = "Object exists, but Legacy_ID__c does not exist."
            elif not field.get("externalId"):
                status = "FIELD_NOT_EXTERNAL_ID"
                detail = "Legacy_ID__c exists but is not marked as an External ID."
            elif not (field.get("createable") or field.get("updateable")):
                status = "FIELD_NOT_WRITABLE"
                detail = "Legacy_ID__c exists but is not writable through the API."
            else:
                status = "READY"
                detail = "Object and usable Legacy_ID__c External ID field exist."
            rows.append({
                "Order": order,
                "Object": object_name,
                "Status": status,
                "ObjectExists": "Yes",
                "LegacyFieldExists": "Yes" if field else "No",
                "FieldType": field.get("type", "") if field else "",
                "Length": field.get("length", "") if field else "",
                "ExternalId": field.get("externalId", "") if field else "",
                "Unique": field.get("unique", "") if field else "",
                "Createable": field.get("createable", "") if field else "",
                "Updateable": field.get("updateable", "") if field else "",
                "ObjectLabel": describe.get("label", ""),
                "Detail": detail,
            })
        except Exception as exc:
            rows.append({
                "Order": order,
                "Object": object_name,
                "Status": "DESCRIBE_ERROR",
                "ObjectExists": "Yes",
                "LegacyFieldExists": "Unknown",
                "FieldType": "",
                "Length": "",
                "ExternalId": "",
                "Unique": "",
                "Createable": "",
                "Updateable": "",
                "ObjectLabel": global_objects[object_name].get("label", ""),
                "Detail": str(exc),
            })

    counts = Counter(row["Status"] for row in rows)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Count"])
    summary.append(["Hierarchy objects audited", len(rows)])
    summary.append(["Objects existing in target", sum(row["ObjectExists"] == "Yes" for row in rows)])
    summary.append(["Objects missing in target", counts["OBJECT_MISSING"]])
    summary.append(["Ready Legacy_ID__c fields", counts["READY"]])
    summary.append(["Legacy_ID__c missing", counts["LEGACY_FIELD_MISSING"]])
    summary.append(["Legacy_ID__c not External ID", counts["FIELD_NOT_EXTERNAL_ID"]])
    summary.append(["Legacy_ID__c not writable", counts["FIELD_NOT_WRITABLE"]])
    summary.append(["Describe errors", counts["DESCRIBE_ERROR"]])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 14

    detail = workbook.create_sheet("Object Audit")
    headers = list(rows[0])
    detail.append(headers)
    for row in rows:
        detail.append([row[column] for column in headers])
    style(detail)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(dict(counts))
    print(args.output)


if __name__ == "__main__":
    main()
