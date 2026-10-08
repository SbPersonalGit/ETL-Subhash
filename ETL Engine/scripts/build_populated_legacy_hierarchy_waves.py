"""
Build UAT→Prod migration hierarchy/waves from Legacy_ID objects.

Rules:
- Scope = objects with Legacy_ID__c populated in UAT (count > 0).
- Only Lookup / Master-Detail (reference) fields are considered.
- A relationship participates in hierarchy ONLY if at least one record has
  a non-null value in that field (SOQL COUNT where field != null).
- Waves = dependency levels from the populated in-scope graph.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import time
from dataclasses import dataclass
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.connector import connector_from_env
from core.object_hierarchy import (
    EXCLUDED_PARENT_OBJECTS,
    EXCLUDED_REFERENCE_FIELDS,
    MappingRow,
    build_hierarchy_analytics,
    build_hierarchy_groups,
    write_object_hierarchy_workbook,
)


DEFAULT_COUNTS_CSV = PROJECT_ROOT / "output" / "legacy_id_soql_counts_updated.csv"
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "uat_to_prod_populated_hierarchy.xlsx"


@dataclass
class RelationshipProbe:
    object_name: str
    field_name: str
    parent_object: str
    relationship_type: str
    is_required: bool
    populated_rows: int | None
    status: str
    detail: str = ""
    parent_in_scope: bool = False


def load_scope_objects(path: Path, min_legacy_count: int = 1) -> list[str]:
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig")))
    objects: list[str] = []
    seen: set[str] = set()
    for row in rows:
        name = (row.get("sObject") or row.get("Object") or "").strip()
        if not name or name in seen:
            continue
        count_raw = row.get("RecordCount") or row.get("LegacyIdPopulatedCount") or "0"
        try:
            count = int(str(count_raw).strip() or "0")
        except ValueError:
            count = 0
        if count < min_legacy_count:
            continue
        seen.add(name)
        objects.append(name)
    return objects


def count_field_populated(connector, object_name: str, field_name: str) -> tuple[int | None, str]:
    soql = f"SELECT COUNT(Id) cnt FROM {object_name} WHERE {field_name} != null"
    try:
        records = connector.query(soql)
        if not records:
            return 0, "ok"
        value = records[0].get("cnt", records[0].get("expr0"))
        return int(value or 0), "ok"
    except Exception as exc:
        return None, str(exc)


def discover_relationships(
    connector,
    object_names: list[str],
) -> tuple[list[RelationshipProbe], list[dict[str, str]], dict[str, dict[str, list[str]]]]:
    scope = set(object_names)
    excluded_parents = {name.lower() for name in EXCLUDED_PARENT_OBJECTS}
    probes: list[RelationshipProbe] = []
    skipped_objects: list[dict[str, str]] = []
    lookup_details: dict[str, dict[str, list[str]]] = {
        name: {"record_lookups": [], "other_lookups": []} for name in object_names
    }

    described: dict[str, dict] = {}
    for object_name in object_names:
        try:
            described[object_name] = connector.describe_object(object_name)
        except Exception as exc:
            skipped_objects.append({"object": object_name, "reason": f"describe_failed: {exc}"})
            continue

    candidates: list[tuple[str, dict, list[str]]] = []
    for object_name, describe in described.items():
        for field in describe.get("fields") or []:
            if field.get("type") != "reference":
                continue
            field_name = (field.get("name") or "").strip()
            if not field_name:
                continue
            reference_targets = [
                (target or "").strip()
                for target in field.get("referenceTo") or []
                if (target or "").strip()
            ]
            is_polymorphic = (
                field.get("polymorphicForeignKey") is True
                or len(reference_targets) > 1
            )

            if field.get("createable") is False:
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {', '.join(reference_targets) or 'unknown'} [non-createable]"
                )
                continue
            if field_name in EXCLUDED_REFERENCE_FIELDS:
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {', '.join(reference_targets) or 'unknown'} [audit/ownership]"
                )
                continue
            if is_polymorphic:
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {', '.join(reference_targets) or 'unknown'} [polymorphic]"
                )
                continue
            if any(target.lower() in excluded_parents for target in reference_targets):
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {', '.join(reference_targets) or 'unknown'} [User dependency]"
                )
                continue
            if not reference_targets:
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → unknown [no target reported]"
                )
                continue

            candidates.append((object_name, field, reference_targets))

    total = len(candidates)
    print(f"Probing populated values on {total} Lookup/Master-Detail fields...")
    for idx, (object_name, field, reference_targets) in enumerate(candidates, 1):
        field_name = field["name"]
        populated, status = count_field_populated(connector, object_name, field_name)
        relationship_type = (
            "Master-Detail" if field.get("cascadeDelete") is True else "Lookup"
        )
        is_required = field.get("nillable") is False
        detail = "" if status == "ok" else status

        for parent_name in reference_targets:
            parent_in_scope = parent_name in scope
            probe = RelationshipProbe(
                object_name=object_name,
                field_name=field_name,
                parent_object=parent_name,
                relationship_type=relationship_type,
                is_required=is_required,
                populated_rows=populated,
                status="OK" if status == "ok" else "ERROR",
                detail=detail,
                parent_in_scope=parent_in_scope,
            )
            probes.append(probe)

            target_text = parent_name
            label = f"{field_name} → {target_text} (populated={populated})"
            if status != "ok":
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {target_text} [count_failed]"
                )
            elif populated and populated > 0 and parent_in_scope:
                lookup_details[object_name]["record_lookups"].append(label)
            elif populated and populated > 0 and not parent_in_scope:
                lookup_details[object_name]["other_lookups"].append(
                    f"{label} [external parent]"
                )
            else:
                lookup_details[object_name]["other_lookups"].append(
                    f"{field_name} → {target_text} [blank]"
                )

        shown = populated if status == "ok" else f"ERROR:{status[:80]}"
        print(f"[{idx}/{total}] {object_name}.{field_name}: {shown}")

    return probes, skipped_objects, lookup_details


def style_sheet(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def write_probe_sheets(
    workbook_path: Path,
    probes: list[RelationshipProbe],
    record_counts: dict[str, int] | None = None,
) -> None:
    """Append population audit sheets onto the hierarchy workbook."""
    from openpyxl import load_workbook

    wb = load_workbook(workbook_path)
    record_counts = record_counts or {}

    def write_rows(title: str, rows: list[RelationshipProbe]) -> None:
        if title in wb.sheetnames:
            del wb[title]
        ws = wb.create_sheet(title)
        ws.append(
            [
                "ObjectName",
                "ParentObject",
                "ParentIdField",
                "RelationshipType",
                "IsRequired",
                "ParentInScope",
                "PopulatedRows",
                "Status",
                "Detail",
            ]
        )
        for row in rows:
            ws.append(
                [
                    row.object_name,
                    row.parent_object,
                    row.field_name,
                    row.relationship_type,
                    "Yes" if row.is_required else "No",
                    "Yes" if row.parent_in_scope else "No",
                    "" if row.populated_rows is None else row.populated_rows,
                    row.status,
                    row.detail,
                ]
            )
        style_sheet(ws)

    internal = [
        p
        for p in probes
        if p.status == "OK" and p.parent_in_scope and (p.populated_rows or 0) > 0
    ]
    blank = [
        p
        for p in probes
        if p.status == "OK" and p.parent_in_scope and (p.populated_rows or 0) == 0
    ]
    external = [
        p
        for p in probes
        if p.status == "OK" and (not p.parent_in_scope) and (p.populated_rows or 0) > 0
    ]
    errors = [p for p in probes if p.status != "OK"]

    write_rows("PopulatedInternalDeps", internal)
    write_rows("BlankRelationships", blank)
    write_rows("PopulatedExternalDeps", external)
    if errors:
        write_rows("CountErrors", errors)

    # Waves summary sheet
    if "Waves" in wb.sheetnames:
        del wb["Waves"]
    waves_ws = wb.create_sheet("Waves", 0)
    hierarchy = wb["ObjectHierarchy"]
    headers = [cell.value for cell in hierarchy[1]]
    wave_idx = headers.index("ParallelBatch")
    obj_idx = headers.index("ObjectName")
    level_idx = headers.index("DependencyLevel")
    parents_idx = headers.index("ParentObjects")
    phase_idx = headers.index("MigrationPhase")

    by_wave: dict[str, list[tuple]] = {}
    for row in hierarchy.iter_rows(min_row=2, values_only=True):
        wave = str(row[wave_idx] or "")
        by_wave.setdefault(wave, []).append(row)

    waves_ws.append(
        [
            "Wave",
            "DependencyLevel",
            "ObjectName",
            "RecordCount",
            "ParentObjects",
            "MigrationPhase",
        ]
    )
    for wave in sorted(by_wave, key=lambda name: (len(name), name)):
        for row in by_wave[wave]:
            object_name = row[obj_idx]
            waves_ws.append(
                [
                    row[wave_idx],
                    row[level_idx],
                    object_name,
                    record_counts.get(str(object_name), ""),
                    row[parents_idx],
                    row[phase_idx],
                ]
            )
    style_sheet(waves_ws)
    wb.save(workbook_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--counts-csv", type=Path, default=DEFAULT_COUNTS_CSV)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--min-legacy-count", type=int, default=1)
    parser.add_argument("--target-profile", default=None)
    args = parser.parse_args()

    load_dotenv(PROJECT_ROOT / ".env", override=True)
    objects = load_scope_objects(args.counts_csv, min_legacy_count=args.min_legacy_count)
    if not objects:
        raise SystemExit(f"No objects found in {args.counts_csv} with count >= {args.min_legacy_count}")

    connector = connector_from_env("TARGET", profile=args.target_profile)
    connector.connect()
    print(f"Connected: {connector.instance_url}")
    print(f"Scope objects (Legacy_ID__c count >= {args.min_legacy_count}): {len(objects)}")

    started = time.time()
    probes, skipped_objects, lookup_details = discover_relationships(connector, objects)

    populated_internal = [
        p
        for p in probes
        if p.status == "OK" and p.parent_in_scope and (p.populated_rows or 0) > 0
    ]
    mappings = [
        MappingRow(
            object_name=p.object_name,
            parent_object=p.parent_object,
            parent_id_field=p.field_name,
            relationship_type=p.relationship_type,
            is_required=p.is_required,
        )
        for p in populated_internal
    ]
    # Deduplicate identical edges
    unique: dict[tuple, MappingRow] = {}
    for row in mappings:
        unique[(row.object_name, row.parent_object, row.parent_id_field)] = row
    mappings = sorted(
        unique.values(),
        key=lambda row: (row.object_name, row.parent_object, row.parent_id_field),
    )

    described_scope = [name for name in objects if name not in {s["object"] for s in skipped_objects}]
    hierarchy_groups = build_hierarchy_groups(mappings, object_names=described_scope)
    hierarchy_analytics = build_hierarchy_analytics(
        mappings, hierarchy_groups, lookup_details=lookup_details
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    try:
        write_object_hierarchy_workbook(
            args.output,
            mappings,
            hierarchy_groups,
            skipped_objects=skipped_objects,
            lookup_details=lookup_details,
            hierarchy_analytics=hierarchy_analytics,
        )
    except PermissionError:
        fallback = args.output.with_name(args.output.stem + "_new.xlsx")
        print(f"Output locked; writing {fallback}")
        write_object_hierarchy_workbook(
            fallback,
            mappings,
            hierarchy_groups,
            skipped_objects=skipped_objects,
            lookup_details=lookup_details,
            hierarchy_analytics=hierarchy_analytics,
        )
        args.output = fallback

    write_probe_sheets(
        args.output,
        probes,
        record_counts={
            (row.get("sObject") or row.get("Object") or "").strip(): int(
                str(row.get("RecordCount") or row.get("LegacyIdPopulatedCount") or "0").strip()
                or "0"
            )
            for row in csv.DictReader(args.counts_csv.open(encoding="utf-8-sig"))
            if (row.get("sObject") or row.get("Object") or "").strip()
        },
    )

    wave_summary: dict[str, list[str]] = {}
    for name, analytics in hierarchy_analytics.items():
        wave = str(analytics.get("parallel_batch") or "Wave ?")
        wave_summary.setdefault(wave, []).append(name)
    for wave in wave_summary:
        wave_summary[wave] = sorted(wave_summary[wave])

    summary = {
        "instance_url": connector.instance_url,
        "scope_object_count": len(objects),
        "described_object_count": len(described_scope),
        "probe_count": len(probes),
        "populated_internal_dependencies": len(populated_internal),
        "blank_internal_relationships": sum(
            1
            for p in probes
            if p.status == "OK" and p.parent_in_scope and (p.populated_rows or 0) == 0
        ),
        "populated_external_dependencies": sum(
            1
            for p in probes
            if p.status == "OK" and (not p.parent_in_scope) and (p.populated_rows or 0) > 0
        ),
        "cycle_group_count": sum(1 for g in hierarchy_groups if g["is_cycle_group"]),
        "wave_count": len(wave_summary),
        "waves": {wave: wave_summary[wave] for wave in sorted(wave_summary, key=lambda w: (len(w), w))},
        "elapsed_seconds": round(time.time() - started, 1),
        "output": str(args.output),
    }
    summary_path = args.output.with_suffix(".summary.json")
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print("\n=== SUMMARY ===")
    print(json.dumps({k: v for k, v in summary.items() if k != "waves"}, indent=2))
    print("\n=== WAVES ===")
    for wave, members in summary["waves"].items():
        print(f"\n{wave} ({len(members)} objects)")
        for member in members:
            parents = hierarchy_analytics[member].get("parent_objects") or []
            parent_text = ", ".join(parents) if parents else "-"
            print(f"  - {member}  [parents: {parent_text}]")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
