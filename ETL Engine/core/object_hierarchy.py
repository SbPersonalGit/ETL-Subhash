from __future__ import annotations

from dataclasses import dataclass
import argparse
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import requests

from core.connector import SFConnector

logger = logging.getLogger(__name__)

EXCLUDED_PARENT_OBJECTS = frozenset({"User"})
EXCLUDED_REFERENCE_FIELDS = frozenset({
    "CreatedById",
    "LastModifiedById",
    "OwnerId",
})


@dataclass(frozen=True)
class MappingRow:
    object_name: str
    parent_object: str
    parent_id_field: str
    relationship_type: str = "Lookup"
    is_required: bool = False


def build_parent_id_mappings(
    connector: SFConnector,
    object_names: list[str],
    skipped_objects: list[dict[str, str]] | None = None,
    lookup_details: dict[str, dict[str, list[str]]] | None = None,
) -> list[MappingRow]:
    scope = _canonicalize_object_names(connector, object_names)
    mappings: list[MappingRow] = []
    described_objects: dict[str, dict] = {}

    for object_name in sorted(scope):
        try:
            described_objects[object_name] = connector.describe_object(object_name)
        except requests.HTTPError as exc:
            if getattr(exc.response, "status_code", None) != 404:
                raise
            reason = "Object is listed globally but has no REST describe resource."
            logger.warning("[ObjectHierarchy] Skipping %s: %s", object_name, reason)
            if skipped_objects is not None:
                skipped_objects.append({"object": object_name, "reason": reason})
            continue

    # Limit relationships to objects with a usable REST describe resource.
    # User-targeting lookups are technical dependencies that otherwise join
    # many unrelated business objects into one large cycle.
    rest_scope = set(described_objects)
    excluded_parents = {name.lower() for name in EXCLUDED_PARENT_OBJECTS}

    for object_name in sorted(rest_scope):
        if lookup_details is not None:
            lookup_details[object_name] = {
                "record_lookups": [],
                "other_lookups": [],
            }
        describe = described_objects[object_name]
        for field in describe.get("fields", []):
            if field.get("type") != "reference":
                continue
            field_name = (field.get("name") or "").strip()
            if not field_name:
                continue
            reference_targets = [
                (target or "").strip()
                for target in field.get("referenceTo") or []
                if (target or "").strip()
            ]
            is_polymorphic = (
                field.get("polymorphicForeignKey") is True
                or len(reference_targets) > 1
            )

            exclusion_reason = ""
            if field.get("createable") is False:
                exclusion_reason = "non-createable"
            elif field_name in EXCLUDED_REFERENCE_FIELDS:
                exclusion_reason = "audit/ownership"
            elif is_polymorphic:
                exclusion_reason = "polymorphic"
            elif any(target.lower() in excluded_parents for target in reference_targets):
                exclusion_reason = "User dependency"
            elif not reference_targets:
                exclusion_reason = "no target reported"
            elif any(target not in rest_scope for target in reference_targets):
                exclusion_reason = "external/not selected"

            target_text = ", ".join(reference_targets) if reference_targets else "unknown"
            detail_text = f"{field_name} → {target_text}"
            if exclusion_reason:
                if lookup_details is not None:
                    lookup_details[object_name]["other_lookups"].append(
                        f"{detail_text} [{exclusion_reason}]"
                    )
                continue

            if lookup_details is not None:
                lookup_details[object_name]["record_lookups"].append(detail_text)
            for parent_name in reference_targets:
                mappings.append(
                    MappingRow(
                        object_name=object_name,
                        parent_object=parent_name,
                        parent_id_field=field_name,
                        relationship_type=(
                            "Master-Detail"
                            if field.get("cascadeDelete") is True
                            else "Lookup"
                        ),
                        is_required=field.get("nillable") is False,
                    )
                )

    mappings.sort(key=lambda row: (row.object_name, row.parent_object, row.parent_id_field))
    return mappings


def build_cycle_groups(mappings: list[MappingRow], object_names: list[str] | None = None) -> list[list[str]]:
    return [
        group["members"]
        for group in build_hierarchy_groups(mappings, object_names=object_names)
        if group["is_cycle_group"]
    ]


def build_hierarchy_groups(
    mappings: list[MappingRow],
    object_names: list[str] | None = None,
) -> list[dict[str, object]]:
    adjacency: dict[str, set[str]] = {}
    all_nodes: set[str] = set()
    self_loops: set[str] = set()

    for object_name in object_names or []:
        normalized = (object_name or "").strip()
        if not normalized:
            continue
        all_nodes.add(normalized)
        adjacency.setdefault(normalized, set())

    for row in mappings:
        all_nodes.add(row.object_name)
        all_nodes.add(row.parent_object)
        adjacency.setdefault(row.object_name, set()).add(row.parent_object)
        adjacency.setdefault(row.parent_object, set())
        if row.object_name == row.parent_object:
            self_loops.add(row.object_name)

    index = 0
    stack: list[str] = []
    stack_members: set[str] = set()
    node_index: dict[str, int] = {}
    lowlink: dict[str, int] = {}
    components: list[list[str]] = []

    def strongconnect(node: str) -> None:
        nonlocal index
        node_index[node] = index
        lowlink[node] = index
        index += 1
        stack.append(node)
        stack_members.add(node)

        for neighbor in adjacency.get(node, ()):
            if neighbor not in node_index:
                strongconnect(neighbor)
                lowlink[node] = min(lowlink[node], lowlink[neighbor])
            elif neighbor in stack_members:
                lowlink[node] = min(lowlink[node], node_index[neighbor])

        if lowlink[node] == node_index[node]:
            component: list[str] = []
            while stack:
                member = stack.pop()
                stack_members.remove(member)
                component.append(member)
                if member == node:
                    break
            components.append(sorted(component))

    for node in sorted(all_nodes):
        if node not in node_index:
            strongconnect(node)

    component_by_member: dict[str, tuple[str, ...]] = {}
    component_nodes: list[tuple[str, ...]] = []
    for component in components:
        ordered = tuple(sorted(component))
        component_nodes.append(ordered)
        for member in ordered:
            component_by_member[member] = ordered

    component_edges: dict[tuple[str, ...], set[tuple[str, ...]]] = {component: set() for component in component_nodes}
    indegree: dict[tuple[str, ...], int] = {component: 0 for component in component_nodes}

    for row in mappings:
        child_group = component_by_member[row.object_name]
        parent_group = component_by_member[row.parent_object]
        if child_group == parent_group:
            continue
        if child_group not in component_edges[parent_group]:
            component_edges[parent_group].add(child_group)
            indegree[child_group] += 1

    ready = sorted(
        [component for component, degree in indegree.items() if degree == 0],
        key=lambda members: (len(members), members),
    )
    dependency_level: dict[tuple[str, ...], int] = {
        component: 1 for component in ready
    }
    ordered_groups: list[tuple[str, ...]] = []
    while ready:
        current = ready.pop(0)
        ordered_groups.append(current)
        for child in sorted(component_edges[current], key=lambda members: (len(members), members)):
            dependency_level[child] = max(
                dependency_level.get(child, 1),
                dependency_level.get(current, 1) + 1,
            )
            indegree[child] -= 1
            if indegree[child] == 0:
                ready.append(child)
                ready.sort(key=lambda members: (len(members), members))

    if len(ordered_groups) != len(component_nodes):
        ordered_groups = sorted(component_nodes, key=lambda members: (len(members), members))

    groups: list[dict[str, object]] = []
    for group_order, members in enumerate(ordered_groups, start=1):
        is_cycle = len(members) > 1 or members[0] in self_loops
        groups.append(
            {
                "group_order": group_order,
                "members": list(members),
                "is_cycle_group": is_cycle,
                "dependency_level": dependency_level.get(members, 1),
            }
        )
    return groups


def build_hierarchy_analytics(
    mappings: list[MappingRow],
    hierarchy_groups: list[dict[str, object]],
    lookup_details: dict[str, dict[str, list[str]]] | None = None,
) -> dict[str, dict[str, object]]:
    group_by_object: dict[str, int] = {}
    group_members: dict[int, set[str]] = {}
    group_levels: dict[int, int] = {}
    cycle_groups: set[int] = set()

    for group in hierarchy_groups:
        group_order = int(group["group_order"])
        members = set(group["members"])
        group_members[group_order] = members
        group_levels[group_order] = int(group.get("dependency_level", 1))
        if group["is_cycle_group"]:
            cycle_groups.add(group_order)
        for member in members:
            group_by_object[member] = group_order

    parents_by_object: dict[str, set[str]] = {
        name: set() for name in group_by_object
    }
    children_by_object: dict[str, set[str]] = {
        name: set() for name in group_by_object
    }
    group_parents: dict[int, set[int]] = {
        group_order: set() for group_order in group_members
    }
    deferred_fields: dict[str, set[str]] = {
        name: set() for name in group_by_object
    }

    for row in mappings:
        child_group = group_by_object[row.object_name]
        parent_group = group_by_object[row.parent_object]
        if child_group == parent_group:
            if child_group in cycle_groups:
                deferred_fields[row.object_name].add(row.parent_id_field)
            continue
        parents_by_object[row.object_name].add(row.parent_object)
        children_by_object[row.parent_object].add(row.object_name)
        group_parents[child_group].add(parent_group)

    transitive_group_parents: dict[int, set[int]] = {}

    def collect_group_parents(group_order: int) -> set[int]:
        if group_order in transitive_group_parents:
            return transitive_group_parents[group_order]
        result: set[int] = set()
        for parent_group in group_parents[group_order]:
            result.add(parent_group)
            result.update(collect_group_parents(parent_group))
        transitive_group_parents[group_order] = result
        return result

    analytics: dict[str, dict[str, object]] = {}
    for object_name, group_order in group_by_object.items():
        parents = sorted(parents_by_object[object_name])
        children = sorted(children_by_object[object_name])
        upstream_groups = collect_group_parents(group_order)
        transitive_parent_count = sum(
            len(group_members[parent_group]) for parent_group in upstream_groups
        )
        other_lookups = (lookup_details or {}).get(object_name, {}).get(
            "other_lookups", []
        )
        polymorphic_count = sum(
            1 for detail in other_lookups if detail.endswith("[polymorphic]")
        )
        deferred_count = len(deferred_fields[object_name]) + polymorphic_count
        level = group_levels[group_order]
        is_cycle = group_order in cycle_groups
        is_foundation = not group_parents[group_order]

        if is_cycle:
            migration_phase = "Cycle Pass 1 + Deferred Lookup Pass 2"
        elif is_foundation:
            migration_phase = "Foundation"
        else:
            migration_phase = "Regular"

        if is_cycle or deferred_count:
            migration_risk = "High"
        elif len(parents) >= 3 or transitive_parent_count >= 5 or len(children) >= 5:
            migration_risk = "Medium"
        else:
            migration_risk = "Low"

        analytics[object_name] = {
            "dependency_level": level,
            "parent_object_count": len(parents),
            "parent_objects": parents,
            "child_object_count": len(children),
            "child_objects": children,
            "transitive_parent_count": transitive_parent_count,
            "critical_path_depth": max(level - 1, 0),
            "is_foundation_object": "Yes" if is_foundation else "No",
            "migration_phase": migration_phase,
            "parallel_batch": f"Wave {level}",
            "deferred_lookup_count": deferred_count,
            "migration_risk": migration_risk,
        }
    return analytics


def write_object_hierarchy_workbook(
    workbook_path: Path,
    mappings: list[MappingRow],
    hierarchy_groups: list[dict[str, object]],
    skipped_objects: list[dict[str, str]] | None = None,
    lookup_details: dict[str, dict[str, list[str]]] | None = None,
    hierarchy_analytics: dict[str, dict[str, object]] | None = None,
) -> Path:
    wb = Workbook()

    mappings_ws = wb.active
    mappings_ws.title = "ParentIdMappings"
    mappings_ws.append([
        "ObjectName",
        "ParentObject",
        "ParentIdField",
        "RelationshipType",
        "IsRequired",
    ])
    for row in mappings:
        mappings_ws.append([
            row.object_name,
            row.parent_object,
            row.parent_id_field,
            row.relationship_type,
            "Yes" if row.is_required else "No",
        ])

    hierarchy_ws = wb.create_sheet("ObjectHierarchy")
    hierarchy_ws.append([
        "Order",
        "GroupOrder",
        "ObjectName",
        "IsCycleGroup",
        "GroupMembers",
        "RecordLookupCount",
        "RecordLookups",
        "OtherLookups",
        "DependencyLevel",
        "ParentObjectCount",
        "ParentObjects",
        "ChildObjectCount",
        "ChildObjects",
        "TransitiveParentCount",
        "CriticalPathDepth",
        "IsFoundationObject",
        "MigrationPhase",
        "ParallelBatch",
        "DeferredLookupCount",
        "MigrationRisk",
    ])
    row_order = 1
    for group in hierarchy_groups:
        members = list(group["members"])
        group_order = int(group["group_order"])
        is_cycle_group = "Yes" if group["is_cycle_group"] else "No"
        member_text = ", ".join(members)
        for object_name in members:
            details = (lookup_details or {}).get(object_name, {})
            analytics = (hierarchy_analytics or {}).get(object_name, {})
            hierarchy_ws.append([
                row_order,
                group_order,
                object_name,
                is_cycle_group,
                member_text,
                len(details.get("record_lookups", [])),
                "\n".join(details.get("record_lookups", [])),
                "\n".join(details.get("other_lookups", [])),
                analytics.get("dependency_level", ""),
                analytics.get("parent_object_count", ""),
                "\n".join(analytics.get("parent_objects", [])),
                analytics.get("child_object_count", ""),
                "\n".join(analytics.get("child_objects", [])),
                analytics.get("transitive_parent_count", ""),
                analytics.get("critical_path_depth", ""),
                analytics.get("is_foundation_object", ""),
                analytics.get("migration_phase", ""),
                analytics.get("parallel_batch", ""),
                analytics.get("deferred_lookup_count", ""),
                analytics.get("migration_risk", ""),
            ])
            row_order += 1

    for ws in (mappings_ws, hierarchy_ws):
        _style_sheet(ws)
    if skipped_objects:
        skipped_ws = wb.create_sheet("SkippedObjects")
        skipped_ws.append(["ObjectName", "Reason"])
        for item in skipped_objects:
            skipped_ws.append([item["object"], item["reason"]])
        _style_sheet(skipped_ws)

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(workbook_path)
    except PermissionError as exc:
        raise PermissionError(
            f"Unable to write workbook '{workbook_path}'. Close the file if it is open in Excel and try again."
        ) from exc
    return workbook_path


def generate_object_hierarchy_workbook(
    connector: SFConnector,
    object_names: list[str],
    workbook_path: Path,
) -> dict[str, object]:
    canonical_scope = _canonicalize_object_names(connector, object_names)
    skipped_objects: list[dict[str, str]] = []
    lookup_details: dict[str, dict[str, list[str]]] = {}
    mappings = build_parent_id_mappings(
        connector,
        list(canonical_scope),
        skipped_objects=skipped_objects,
        lookup_details=lookup_details,
    )
    skipped_names = {item["object"] for item in skipped_objects}
    described_scope = canonical_scope - skipped_names
    hierarchy_groups = build_hierarchy_groups(
        mappings, object_names=sorted(described_scope)
    )
    hierarchy_analytics = build_hierarchy_analytics(
        mappings, hierarchy_groups, lookup_details=lookup_details
    )
    write_object_hierarchy_workbook(
        workbook_path,
        mappings,
        hierarchy_groups,
        skipped_objects=skipped_objects,
        lookup_details=lookup_details,
        hierarchy_analytics=hierarchy_analytics,
    )
    sheet_names = ["ParentIdMappings", "ObjectHierarchy"]
    if skipped_objects:
        sheet_names.append("SkippedObjects")
    return {
        "path": str(workbook_path),
        "object_count": len(described_scope),
        "requested_object_count": len(canonical_scope),
        "skipped_object_count": len(skipped_objects),
        "skipped_objects": skipped_objects,
        "mapping_count": len(mappings),
        "cycle_group_count": sum(1 for group in hierarchy_groups if group["is_cycle_group"]),
        "cycle_object_count": sum(len(group["members"]) for group in hierarchy_groups if group["is_cycle_group"]),
        "sheet_names": sheet_names,
    }


def _canonicalize_object_names(connector: SFConnector, object_names: list[str]) -> set[str]:
    requested = []
    seen_requested: set[str] = set()
    for name in object_names:
        normalized = (name or "").strip()
        if not normalized:
            continue
        lowered = normalized.lower()
        if lowered in seen_requested:
            continue
        seen_requested.add(lowered)
        requested.append(normalized)

    if hasattr(connector, "describe_global"):
        available = {
            (entry.get("name") or "").strip().lower(): (entry.get("name") or "").strip()
            for entry in connector.describe_global()
            if (entry.get("name") or "").strip()
        }
    else:
        available = {}

    canonical_scope: set[str] = set()
    for name in requested:
        canonical_scope.add(available.get(name.lower(), name))
    return canonical_scope


def _style_sheet(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = wrap

    for column in ws.columns:
        values = [len(str(cell.value or "")) for cell in column]
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(max(max(values, default=10) + 2, 14), 48)


def _read_object_scope(path: Path) -> list[str]:
    objects: list[str] = []
    seen: set[str] = set()
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        name = raw_line.strip()
        if not name or name.startswith("#") or name in seen:
            continue
        seen.add(name)
        objects.append(name)
    return objects


def _parse_cli_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate an object hierarchy workbook directly from Salesforce metadata."
    )
    parser.add_argument(
        "--org",
        choices=["source", "target"],
        default="target",
        help="Salesforce org role to inspect.",
    )
    parser.add_argument(
        "--profile",
        default="default",
        help="Configured SOURCE/TARGET environment profile.",
    )
    parser.add_argument(
        "--objects-file",
        type=Path,
        default=Path("config/waygate_object_scope.txt"),
        help="Text file containing one Salesforce object API name per line.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("ui/waygate_object_hierarchy.xlsx"),
        help="Output XLSX path.",
    )
    return parser.parse_args()


def main() -> int:
    from dotenv import load_dotenv
    from core.connector import connector_from_env

    args = _parse_cli_args()
    project_root = Path(__file__).resolve().parents[1]
    load_dotenv(project_root / ".env", override=True)

    objects_file = args.objects_file.resolve()
    if not objects_file.is_file():
        raise SystemExit(f"Object scope file not found: {objects_file}")
    object_names = _read_object_scope(objects_file)
    if not object_names:
        raise SystemExit(f"No object API names found in: {objects_file}")

    connector = connector_from_env(args.org.upper(), args.profile)
    login_url = connector.login_url.lower()
    if (
        "your-source-org" in login_url
        or "your-target-org" in login_url
        or "example." in login_url
    ):
        profile_suffix = "" if args.profile.strip().lower() == "default" else f"_{args.profile.strip().upper()}"
        env_key = f"{args.org.upper()}{profile_suffix}_SF_LOGIN_URL"
        raise SystemExit(
            f"{args.org.title()} profile '{args.profile}' is not configured: "
            f"{env_key} contains the placeholder URL '{connector.login_url}'. "
            f"Set a real Salesforce login URL in .env or run with --org target."
        )
    try:
        connector.connect()
    except requests.RequestException as exc:
        raise SystemExit(
            f"Unable to connect to {args.org} profile '{args.profile}' "
            f"at {connector.login_url}: {exc}"
        ) from exc
    result = generate_object_hierarchy_workbook(
        connector,
        object_names,
        args.output.resolve(),
    )
    print(f"Org: {connector.instance_url}")
    print(f"Objects requested: {result['requested_object_count']}")
    print(f"Objects written: {result['object_count']}")
    print(f"Relationships: {result['mapping_count']}")
    print(f"Skipped objects: {result['skipped_object_count']}")
    for item in result["skipped_objects"]:
        print(f"  - {item['object']}: {item['reason']}")
    print(f"Workbook: {result['path']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
